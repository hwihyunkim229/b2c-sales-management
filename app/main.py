from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from uuid import uuid4
from urllib.parse import quote
import io
import math
import re

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill

from fastapi import FastAPI, File, Form, Query, Request, UploadFile
from fastapi.responses import JSONResponse, RedirectResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from app.database import get_connection, initialize_database
from app.services.excel_importer import import_workbook
from app.services.excel_importer import SUPPORTED_PLATFORMS
from app.services.data_validator import build_validation_report
from app.security import (
    INITIAL_CREDENTIALS_PATH,
    SESSION_COOKIE_NAME,
    SESSION_DURATION_SECONDS,
    create_session,
    delete_session,
    ensure_initial_admin,
    get_session_user,
    hash_password,
    password_error,
    remove_initial_credentials,
    verify_password,
)


APP_DIRECTORY = Path(__file__).resolve().parent
PROJECT_ROOT = APP_DIRECTORY.parent
UPLOAD_DIRECTORY = PROJECT_ROOT / "uploads"


def ensure_batch_source_file(connection, batch) -> Path:
    stored_path = UPLOAD_DIRECTORY / batch["stored_filename"]
    if stored_path.exists():
        return stored_path

    source_content = batch["source_content"]
    if source_content is None:
        source_row = connection.execute(
            "SELECT source_content FROM import_batches WHERE id = ?",
            (batch["id"],),
        ).fetchone()
        source_content = source_row["source_content"] if source_row else None

    if source_content is None:
        raise FileNotFoundError(
            f"업로드 원본을 찾을 수 없습니다: {batch['original_filename']}"
        )

    UPLOAD_DIRECTORY.mkdir(parents=True, exist_ok=True)
    stored_path.write_bytes(bytes(source_content))
    return stored_path


def restore_completed_source_files(connection) -> None:
    batches = connection.execute(
        """
        SELECT id, original_filename, stored_filename, source_content
        FROM import_batches
        WHERE status = 'completed'
        """
    ).fetchall()
    for batch in batches:
        ensure_batch_source_file(connection, batch)


def backfill_source_content(connection) -> int:
    batches = connection.execute(
        """
        SELECT id, stored_filename
        FROM import_batches
        WHERE source_content IS NULL
        """
    ).fetchall()
    updated = 0
    for batch in batches:
        file_path = UPLOAD_DIRECTORY / batch["stored_filename"]
        if not file_path.exists():
            continue
        connection.execute(
            "UPDATE import_batches SET source_content = ? WHERE id = ?",
            (file_path.read_bytes(), batch["id"]),
        )
        updated += 1
    return updated


@asynccontextmanager
async def lifespan(app: FastAPI):
    initialize_database()
    UPLOAD_DIRECTORY.mkdir(parents=True, exist_ok=True)
    with get_connection() as connection:
        created = ensure_initial_admin(connection)
        restored_sources = backfill_source_content(connection)
    if created:
        if INITIAL_CREDENTIALS_PATH.exists():
            print(f"[B2C Sales] 최초 관리자 자격증명: {INITIAL_CREDENTIALS_PATH}")
        else:
            print("[B2C Sales] 환경변수로 최초 관리자 계정을 생성했습니다.")
    if restored_sources:
        print(f"[B2C Sales] 기존 업로드 원본 {restored_sources}개를 DB에 저장했습니다.")
    yield


app = FastAPI(
    title="B2C Sales Management",
    lifespan=lifespan,
)
app.mount(
    "/static",
    StaticFiles(directory=APP_DIRECTORY / "static"),
    name="static",
)

templates = Jinja2Templates(
    directory=APP_DIRECTORY / "templates",
)


PUBLIC_PATHS = {"/login", "/health", "/favicon.ico"}


@app.middleware("http")
async def authentication_middleware(request: Request, call_next):
    path = request.url.path
    token = request.cookies.get(SESSION_COOKIE_NAME)
    user = None
    if token:
        with get_connection() as connection:
            user = get_session_user(connection, token)

    request.state.user = user

    if path.startswith("/static/") or path in PUBLIC_PATHS:
        return await call_next(request)

    if user is None:
        next_path = path if request.method == "GET" else "/"
        return RedirectResponse(
            f"/login?next={quote(next_path, safe='/')}",
            status_code=303,
        )

    if user["must_change_password"] and path not in {
        "/change-password",
        "/logout",
    }:
        return RedirectResponse("/change-password", status_code=303)

    if path.startswith("/users") and user["role"] != "admin":
        return templates.TemplateResponse(
            request=request,
            name="permission_denied.html",
            context={},
            status_code=403,
        )

    if (
        user["role"] == "viewer"
        and path.startswith("/sales/")
        and path.endswith("/edit")
    ):
        return templates.TemplateResponse(
            request=request,
            name="permission_denied.html",
            context={},
            status_code=403,
        )

    if request.method not in {"GET", "HEAD", "OPTIONS"} and path not in {
        "/change-password",
        "/logout",
    }:
        role = user["role"]
        reviewer_restricted = path.startswith("/imports") or path.startswith("/hometax")
        if role == "viewer" or (role == "reviewer" and reviewer_restricted):
            return templates.TemplateResponse(
                request=request,
                name="permission_denied.html",
                context={},
                status_code=403,
            )

    return await call_next(request)


def set_session_cookie(response: RedirectResponse, token: str, request: Request) -> None:
    response.set_cookie(
        SESSION_COOKIE_NAME,
        token,
        max_age=SESSION_DURATION_SECONDS,
        httponly=True,
        secure=request.url.scheme == "https",
        samesite="lax",
        path="/",
    )


def safe_next_path(value: str) -> str:
    if value.startswith("/") and not value.startswith("//"):
        return value
    return "/"


@app.api_route("/health", methods=["GET", "HEAD"])
async def health(request: Request):
    if request.method == "HEAD":
        return Response(status_code=200)
    return JSONResponse({"status": "ok"})


@app.get("/login")
async def login_page(request: Request, next: str = "/", error: str = ""):
    if request.state.user:
        destination = (
            "/change-password"
            if request.state.user["must_change_password"]
            else safe_next_path(next)
        )
        return RedirectResponse(destination, status_code=303)
    return templates.TemplateResponse(
        request=request,
        name="login.html",
        context={"next": safe_next_path(next), "error": error},
    )


@app.post("/login")
async def login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    next: str = Form("/"),
):
    with get_connection() as connection:
        connection.execute(
            "DELETE FROM auth_sessions WHERE expires_at <= ?",
            (datetime.now(timezone.utc).isoformat(timespec="seconds"),),
        )
        user = connection.execute(
            "SELECT * FROM users WHERE username = ? AND is_active = 1",
            (username.strip(),),
        ).fetchone()
        valid = bool(
            user
            and verify_password(
                password,
                user["password_hash"],
                user["password_salt"],
                user["password_iterations"],
            )
        )
        if not valid:
            connection.execute(
                """
                INSERT INTO audit_logs (action, target_type, target_id, description)
                VALUES ('login_failed', 'user', ?, ?)
                """,
                (user["id"] if user else None, f"로그인 실패: {username.strip()}"),
            )
            return RedirectResponse(
                f"/login?error=invalid&next={quote(safe_next_path(next), safe='/')}",
                status_code=303,
            )

        token = create_session(connection, user["id"])
        connection.execute(
            "UPDATE users SET last_login_at = CURRENT_TIMESTAMP WHERE id = ?",
            (user["id"],),
        )
        connection.execute(
            """
            INSERT INTO audit_logs (action, target_type, target_id, description)
            VALUES ('login', 'user', ?, ?)
            """,
            (user["id"], f"{user['username']} 로그인"),
        )

    destination = "/change-password" if user["must_change_password"] else safe_next_path(next)
    response = RedirectResponse(destination, status_code=303)
    set_session_cookie(response, token, request)
    return response


@app.get("/change-password")
async def change_password_page(request: Request, error: str = ""):
    return templates.TemplateResponse(
        request=request,
        name="change_password.html",
        context={"error": error},
    )


@app.post("/change-password")
async def change_password(
    request: Request,
    current_password: str = Form(...),
    new_password: str = Form(...),
    confirm_password: str = Form(...),
):
    user = request.state.user
    if new_password != confirm_password:
        return RedirectResponse("/change-password?error=confirm", status_code=303)
    validation_error = password_error(new_password)
    if validation_error:
        return templates.TemplateResponse(
            request=request,
            name="change_password.html",
            context={"error": "policy", "policy_error": validation_error},
            status_code=400,
        )
    if verify_password(
        new_password,
        user["password_hash"],
        user["password_salt"],
        user["password_iterations"],
    ):
        return RedirectResponse("/change-password?error=same", status_code=303)
    if not verify_password(
        current_password,
        user["password_hash"],
        user["password_salt"],
        user["password_iterations"],
    ):
        return RedirectResponse("/change-password?error=current", status_code=303)

    password_hash, password_salt, iterations = hash_password(new_password)
    with get_connection() as connection:
        connection.execute(
            """
            UPDATE users
            SET password_hash = ?, password_salt = ?, password_iterations = ?,
                must_change_password = 0, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (password_hash, password_salt, iterations, user["id"]),
        )
        connection.execute("DELETE FROM auth_sessions WHERE user_id = ?", (user["id"],))
        token = create_session(connection, user["id"])
        connection.execute(
            """
            INSERT INTO audit_logs (action, target_type, target_id, description)
            VALUES ('password_changed', 'user', ?, ?)
            """,
            (user["id"], f"{user['username']} 비밀번호 변경"),
        )

    remove_initial_credentials()
    response = RedirectResponse("/", status_code=303)
    set_session_cookie(response, token, request)
    return response


@app.post("/logout")
async def logout(request: Request):
    token = request.cookies.get(SESSION_COOKIE_NAME)
    user = request.state.user
    with get_connection() as connection:
        delete_session(connection, token)
        connection.execute(
            """
            INSERT INTO audit_logs (action, target_type, target_id, description)
            VALUES ('logout', 'user', ?, ?)
            """,
            (user["id"], f"{user['username']} 로그아웃"),
        )
    response = RedirectResponse("/login", status_code=303)
    response.delete_cookie(SESSION_COOKIE_NAME, path="/")
    return response


ROLE_LABELS = {
    "admin": "관리자",
    "reviewer": "검수자",
    "viewer": "조회자",
}


@app.get("/users")
async def users_page(request: Request, message: str = "", error: str = ""):
    with get_connection() as connection:
        users = connection.execute(
            """
            SELECT id, username, display_name, role, must_change_password,
                   is_active, last_login_at, created_at
            FROM users
            ORDER BY is_active DESC, id
            """
        ).fetchall()
    return templates.TemplateResponse(
        request=request,
        name="users.html",
        context={
            "users": users,
            "role_labels": ROLE_LABELS,
            "message": message,
            "error": error,
        },
    )


@app.post("/users")
async def create_user(
    request: Request,
    username: str = Form(...),
    display_name: str = Form(...),
    role: str = Form(...),
    temporary_password: str = Form(...),
    confirm_password: str = Form(...),
):
    username = username.strip().lower()
    display_name = display_name.strip()
    if not re.fullmatch(r"[a-z0-9._-]{3,30}", username):
        return RedirectResponse("/users?error=username", status_code=303)
    if not display_name or len(display_name) > 30:
        return RedirectResponse("/users?error=name", status_code=303)
    if role not in ROLE_LABELS:
        return RedirectResponse("/users?error=role", status_code=303)
    if temporary_password != confirm_password:
        return RedirectResponse("/users?error=confirm", status_code=303)
    if password_error(temporary_password):
        return RedirectResponse("/users?error=password", status_code=303)

    password_hash, password_salt, iterations = hash_password(temporary_password)
    with get_connection() as connection:
        if connection.execute(
            "SELECT 1 FROM users WHERE username = ?", (username,)
        ).fetchone():
            return RedirectResponse("/users?error=duplicate", status_code=303)
        cursor = connection.execute(
            """
            INSERT INTO users (
                username, display_name, role, password_hash, password_salt,
                password_iterations, must_change_password, is_active
            ) VALUES (?, ?, ?, ?, ?, ?, 1, 1)
            RETURNING id
            """,
            (
                username,
                display_name,
                role,
                password_hash,
                password_salt,
                iterations,
            ),
        )
        created_user = cursor.fetchone()
        connection.execute(
            """
            INSERT INTO audit_logs (action, target_type, target_id, description)
            VALUES ('user_created', 'user', ?, ?)
            """,
            (
                created_user["id"],
                f"{request.state.user['username']} 계정 생성: {username} ({ROLE_LABELS[role]})",
            ),
        )
    return RedirectResponse("/users?message=created", status_code=303)


@app.post("/users/{user_id}/role")
async def change_user_role(request: Request, user_id: int, role: str = Form(...)):
    if role not in ROLE_LABELS:
        return RedirectResponse("/users?error=role", status_code=303)
    if user_id == request.state.user["id"]:
        return RedirectResponse("/users?error=self_role", status_code=303)
    with get_connection() as connection:
        target = connection.execute(
            "SELECT username FROM users WHERE id = ?", (user_id,)
        ).fetchone()
        if not target:
            return RedirectResponse("/users?error=missing", status_code=303)
        connection.execute(
            "UPDATE users SET role = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (role, user_id),
        )
        connection.execute(
            """
            INSERT INTO audit_logs (action, target_type, target_id, description)
            VALUES ('user_role_changed', 'user', ?, ?)
            """,
            (user_id, f"{target['username']} 권한을 {ROLE_LABELS[role]}로 변경"),
        )
    return RedirectResponse("/users?message=role", status_code=303)


@app.post("/users/{user_id}/status")
async def toggle_user_status(request: Request, user_id: int):
    if user_id == request.state.user["id"]:
        return RedirectResponse("/users?error=self_status", status_code=303)
    with get_connection() as connection:
        target = connection.execute(
            "SELECT username, is_active FROM users WHERE id = ?", (user_id,)
        ).fetchone()
        if not target:
            return RedirectResponse("/users?error=missing", status_code=303)
        new_status = 0 if target["is_active"] else 1
        connection.execute(
            "UPDATE users SET is_active = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (new_status, user_id),
        )
        if not new_status:
            connection.execute("DELETE FROM auth_sessions WHERE user_id = ?", (user_id,))
        connection.execute(
            """
            INSERT INTO audit_logs (action, target_type, target_id, description)
            VALUES ('user_status_changed', 'user', ?, ?)
            """,
            (user_id, f"{target['username']} 계정 {'활성화' if new_status else '비활성화'}"),
        )
    return RedirectResponse("/users?message=status", status_code=303)


@app.post("/users/{user_id}/reset-password")
async def reset_user_password(
    request: Request,
    user_id: int,
    temporary_password: str = Form(...),
    confirm_password: str = Form(...),
):
    if temporary_password != confirm_password:
        return RedirectResponse("/users?error=reset_confirm", status_code=303)
    if password_error(temporary_password):
        return RedirectResponse("/users?error=reset_password", status_code=303)
    password_hash, password_salt, iterations = hash_password(temporary_password)
    with get_connection() as connection:
        target = connection.execute(
            "SELECT username FROM users WHERE id = ?", (user_id,)
        ).fetchone()
        if not target:
            return RedirectResponse("/users?error=missing", status_code=303)
        connection.execute(
            """
            UPDATE users
            SET password_hash = ?, password_salt = ?, password_iterations = ?,
                must_change_password = 1, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (password_hash, password_salt, iterations, user_id),
        )
        connection.execute("DELETE FROM auth_sessions WHERE user_id = ?", (user_id,))
        connection.execute(
            """
            INSERT INTO audit_logs (action, target_type, target_id, description)
            VALUES ('password_reset', 'user', ?, ?)
            """,
            (user_id, f"{target['username']} 임시 비밀번호 재설정"),
        )
    return RedirectResponse("/users?message=reset", status_code=303)


def year_month(value: object) -> str:
    if value in (None, ""):
        return "-"

    text_value = str(value).strip()

    if len(text_value) >= 7:
        return text_value[:7]

    return text_value


def date_only(value: object) -> str:
    if value in (None, ""):
        return "-"
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip().split("T", 1)[0].split(" ", 1)[0]


KOREA_TIMEZONE = timezone(timedelta(hours=9))


def korean_datetime(value: object) -> str:
    if value in (None, ""):
        return "-"

    if isinstance(value, datetime):
        parsed_value = value
    else:
        try:
            parsed_value = datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
        except ValueError:
            return str(value)

    if parsed_value.tzinfo is None:
        parsed_value = parsed_value.replace(tzinfo=timezone.utc)

    return parsed_value.astimezone(KOREA_TIMEZONE).strftime("%Y-%m-%d %H:%M")


def optional_number(
    value: str,
    minimum: int,
    maximum: int,
) -> int | None:
    if not value.strip():
        return None

    try:
        parsed_value = int(value)
    except ValueError:
        return None

    if minimum <= parsed_value <= maximum:
        return parsed_value

    return None


def excel_date_value(value: object) -> object:
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    try:
        return datetime.fromisoformat(str(value).strip()).date()
    except ValueError:
        return value


def add_export_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[object]],
    currency_columns: tuple[int, ...] = (),
    date_columns: tuple[int, ...] = (),
) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append(headers)

    for row in rows:
        worksheet.append(row)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B5B")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 30

    for column in currency_columns:
        for cell in worksheet.iter_cols(
            min_col=column, max_col=column, min_row=2
        ):
            for item in cell:
                item.number_format = "#,##0"

    for column in date_columns:
        for cell in worksheet.iter_cols(
            min_col=column, max_col=column, min_row=2
        ):
            for item in cell:
                item.value = excel_date_value(item.value)
                item.number_format = "yyyy-mm-dd"

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

    for column in worksheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 36)
        worksheet.column_dimensions[column[0].column_letter].width = width


def dated_excel_filename(korean_name: str) -> str:
    today = datetime.now().strftime("%Y%m%d")
    return f"{korean_name}_{today}.xlsx"


def workbook_response(workbook: Workbook, filename: str) -> StreamingResponse:
    if "Sheet" in workbook.sheetnames and len(workbook.sheetnames) > 1:
        del workbook["Sheet"]

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f"attachment; filename=download.xlsx; "
                f"filename*=UTF-8''{quote(filename)}"
            )
        },
    )

templates.env.filters["year_month"] = year_month
templates.env.filters["date_only"] = date_only
templates.env.filters["korean_datetime"] = korean_datetime


SALES_EXPRESSION = """
    credit_card_amount
    + cash_receipt_amount
    + tax_invoice_amount
    + other_amount
"""

SALE_QUANTITY_EXPRESSION = """
    CASE
        WHEN COALESCE(item_type, '') IN ('제품취소', '취소', '본품취소')
         AND REPLACE(COALESCE(product_option, ''), ' ', '') NOT IN (
             '오므론', '배송비', '사이즈미정'
         )
        THEN -ABS(COALESCE(quantity, 0))
        WHEN COALESCE(item_type, '') IN ('제품', '상품', '제품확정')
         AND REPLACE(COALESCE(product_option, ''), ' ', '') NOT IN (
             '오므론', '배송비', '사이즈미정'
         )
        THEN COALESCE(quantity, 0)
        ELSE 0
    END
"""

PRODUCT_TYPE_EQUIVALENTS = {
    "제품": ("제품", "상품"),
    "제품취소": ("제품취소", "취소", "본품취소"),
    "제품확정": ("제품확정",),
    "부가제품": ("부가제품", "부가상품"),
    "배송비": ("배송비",),
    "사이즈미정": ("사이즈미정",),
    "수수료": ("수수료",),
}


def normalized_product_type(value: object) -> str:
    label = str(value or "").strip()
    for standard, equivalents in PRODUCT_TYPE_EQUIVALENTS.items():
        if label in equivalents:
            return standard
    return label or "제품"


def add_product_type_condition(
    conditions: list[str], parameters: list[object], value: str
) -> None:
    equivalents = PRODUCT_TYPE_EQUIVALENTS.get(value)
    if equivalents is None:
        conditions.append("item_type = ?")
        parameters.append(value)
        return
    placeholders = ", ".join("?" for _ in equivalents)
    conditions.append(f"item_type IN ({placeholders})")
    parameters.extend(equivalents)

PRODUCT_CODES = {
    "rev1_7": "SL-P-PD-00033",
    "rev1_8": "SL-P-PD-00034",
    "rev1_9": "SL-P-PD-00035",
    "rev1_10": "SL-P-PD-00036",
    "rev1_11": "SL-P-PD-00037",
    "rev1_12": "SL-P-PD-00038",
    "rev1_13": "SL-P-PD-00039",
    "rev2_7": "SL-P-PD-00054",
    "rev2_8": "SL-P-PD-00055",
    "rev2_9": "SL-P-PD-00056",
    "rev2_10": "SL-P-PD-00057",
    "rev2_11": "SL-P-PD-00058",
    "rev2_12": "SL-P-PD-00059",
    "rev2_13": "SL-P-PD-00060",
}


def normalized_size_key(value: object) -> str:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        numeric_value = float(value)
        if numeric_value.is_integer():
            return f"rev1_{int(numeric_value)}"

    label = str(value or "").strip().lower().replace(" ", "").replace(".", "")
    label = label.replace("-", "_")
    numeric_match = re.fullmatch(r"(\d+)\.0+", str(value or "").strip())
    if numeric_match:
        return f"rev1_{numeric_match.group(1)}"
    if re.fullmatch(r"\d+", label):
        return f"rev1_{label}"
    match = re.fullmatch(r"rev([12])_?(\d+)", label)
    if match:
        return f"rev{match.group(1)}_{match.group(2)}"
    return label


def product_code_for_size(value: object, item_type: object = None) -> str:
    original_value = str(value or "").strip()
    mapped_code = PRODUCT_CODES.get(normalized_size_key(value))
    if mapped_code is not None:
        return mapped_code
    normalized_option = original_value.replace(" ", "")
    if normalized_option in {"오므론", "배송비", "사이즈미정"}:
        return original_value
    if normalized_product_type(item_type) == "제품취소":
        return "제품취소"
    return original_value or "-"


def product_type_label(value: object) -> str:
    return normalized_product_type(value)


def dashboard_size_label(value: object) -> str:
    label = str(value or "미지정").strip()
    normalized_label = normalized_size_key(value)
    rev1_match = re.fullmatch(r"rev1_(\d+)", normalized_label)
    if rev1_match:
        return f"Rev.1_{rev1_match.group(1)}"
    rev2_match = re.fullmatch(r"rev2_(\d+)", normalized_label)
    if rev2_match:
        return f"Rev.2_{rev2_match.group(1)}"
    return label


def dashboard_size_sort_key(value: object) -> tuple[object, ...]:
    label = str(value or "미지정").strip()
    normalized_label = normalized_size_key(value)
    rev1_match = re.fullmatch(r"rev1_(\d+)", normalized_label)
    if rev1_match:
        size = rev1_match.group(1)
        return (0, float(size), size)
    rev2_match = re.fullmatch(r"rev2_(\d+)", normalized_label)
    if rev2_match:
        size = rev2_match.group(1)
        return (1, float(size), size)
    return (2, math.inf, label)


templates.env.filters["product_code"] = product_code_for_size
templates.env.filters["product_type"] = product_type_label
templates.env.filters["product_size"] = dashboard_size_label


@app.get("/")
def dashboard(
    request: Request,
    size_month: str = Query(default=""),
):
    selected_size_month = optional_number(size_month, 1, 12)
    with get_connection() as connection:
        summary = connection.execute(
            """
            SELECT
                COUNT(*) AS record_count,
                COALESCE(SUM({quantity_expression}), 0) AS quantity,
                COALESCE(
                    SUM({sales_expression}),
                    0
                ) AS sales_amount,
                COALESCE(SUM(settlement_amount), 0) AS settlement_amount,
                COALESCE(SUM(fee_amount_1 + fee_amount_2), 0) AS fee_amount
            FROM sales_records
            WHERE is_active = 1 AND batch_is_active = 1
            """.format(
                quantity_expression=SALE_QUANTITY_EXPRESSION,
                sales_expression=SALES_EXPRESSION,
            )
        ).fetchone()

        platforms = connection.execute(
            """
            SELECT
                platform,
                COUNT(*) AS record_count,
                COALESCE(SUM({quantity_expression}), 0) AS quantity,
                COALESCE(
                    SUM({sales_expression}),
                    0
                ) AS sales_amount,
                COALESCE(SUM(settlement_amount), 0) AS settlement_amount
            FROM sales_records
            WHERE is_active = 1 AND batch_is_active = 1
            GROUP BY platform
            ORDER BY sales_amount DESC
            """.format(
                quantity_expression=SALE_QUANTITY_EXPRESSION,
                sales_expression=SALES_EXPRESSION,
            )
        ).fetchall()

        monthly_sales = connection.execute(
            """
            SELECT
                SUBSTR(tax_date, 1, 7) AS period,
                COALESCE(
                    SUM({sales_expression}),
                    0
                ) AS sales_amount
            FROM sales_records
            WHERE tax_date IS NOT NULL AND is_active = 1 AND batch_is_active = 1
            GROUP BY period
            ORDER BY period
            """.format(
                sales_expression=SALES_EXPRESSION,
            )
        ).fetchall()

        size_conditions = ["is_active = 1", "batch_is_active = 1"]
        size_parameters = []
        if selected_size_month is not None:
            size_conditions.append("sales_month = ?")
            size_parameters.append(selected_size_month)
        size_where_clause = " AND ".join(size_conditions)

        platform_sizes = connection.execute(
            f"""
            SELECT
                platform,
                COALESCE(product_option, '미지정') AS product_option,
                CASE
                    WHEN COALESCE(item_type, '') IN (
                        '제품취소', '취소', '본품취소'
                    ) THEN 'cancel'
                    ELSE 'sale'
                END AS quantity_kind,
                COALESCE(SUM({SALE_QUANTITY_EXPRESSION}), 0) AS quantity
            FROM sales_records
            WHERE {size_where_clause}
            GROUP BY platform, product_option, quantity_kind
            HAVING ABS(SUM({SALE_QUANTITY_EXPRESSION})) > 0.001
            ORDER BY platform, product_option, quantity_kind DESC
            """,
            size_parameters,
        ).fetchall()

        size_months = connection.execute(
            """
            SELECT DISTINCT sales_month
            FROM sales_records
            WHERE is_active = 1 AND batch_is_active = 1
              AND sales_month IS NOT NULL
            ORDER BY sales_month
            """
        ).fetchall()

    platform_size_groups = []
    platform_size_lookup = {}
    for row in platform_sizes:
        platform_name = row["platform"]
        if platform_name not in platform_size_lookup:
            group = {
                "platform": platform_name,
                "items": [],
                "item_lookup": {},
                "total_quantity": 0,
                "sale_quantity": 0,
                "cancel_quantity": 0,
            }
            platform_size_lookup[platform_name] = group
            platform_size_groups.append(group)
        group = platform_size_lookup[platform_name]
        quantity = row["quantity"] or 0
        display_label = product_code_for_size(row["product_option"])
        item_key = (display_label, row["quantity_kind"])
        if item_key not in group["item_lookup"]:
            item = {
                "product_option": row["product_option"],
                "display_label": display_label,
                "quantity_kind": row["quantity_kind"],
                "quantity_kind_label": (
                    "제품취소" if row["quantity_kind"] == "cancel"
                    else "제품·제품확정"
                ),
                "quantity": 0,
                "absolute_quantity": 0,
            }
            group["item_lookup"][item_key] = item
            group["items"].append(item)
        item = group["item_lookup"][item_key]
        item["quantity"] += quantity
        item["absolute_quantity"] = abs(item["quantity"])
        group["total_quantity"] += quantity
        if row["quantity_kind"] == "cancel":
            group["cancel_quantity"] += quantity
        else:
            group["sale_quantity"] += quantity
    for group in platform_size_groups:
        group.pop("item_lookup", None)
        group["items"].sort(
            key=lambda item: (
                dashboard_size_sort_key(item["product_option"]),
                1 if item["quantity_kind"] == "cancel" else 0,
            )
        )
        absolute_total = sum(
            item["absolute_quantity"] for item in group["items"]
        )
        for item in group["items"]:
            item["share"] = (
                item["absolute_quantity"] / absolute_total * 100
                if absolute_total else 0
            )

    maximum_monthly_sales = max(
        (row["sales_amount"] for row in monthly_sales),
        default=0,
    )
    maximum_platform_sales = max(
        (row["sales_amount"] for row in platforms),
        default=0,
    )

    return templates.TemplateResponse(
        request=request,
        name="dashboard.html",
        context={
            "summary": summary,
            "platforms": platforms,
            "monthly_sales": monthly_sales,
            "platform_sizes": platform_sizes,
            "platform_size_groups": platform_size_groups,
            "size_months": size_months,
            "selected_size_month": selected_size_month,
            "maximum_monthly_sales": maximum_monthly_sales,
            "maximum_platform_sales": maximum_platform_sales,
        },
    )


@app.get("/validation")
def validation_center(request: Request):
    with get_connection() as connection:
        restore_completed_source_files(connection)
        report = build_validation_report(
            connection=connection,
            upload_directory=UPLOAD_DIRECTORY,
        )

    return templates.TemplateResponse(
        request=request,
        name="validation.html",
        context={"report": report},
    )


@app.get("/exports/validation.xlsx")
def export_validation():
    with get_connection() as connection:
        restore_completed_source_files(connection)
        report = build_validation_report(connection, UPLOAD_DIRECTORY)

    workbook = Workbook()
    add_export_sheet(workbook, "플랫폼 대사",
        ["플랫폼", "원본 건수", "웹 건수", "건수 차이", "원본 수량", "웹 수량", "수량 차이", "원본 매출", "웹 매출", "매출 차이"],
        [[r["platform"], r["source_count"], r["database_count"], r["count_difference"], r["source_quantity"], r["database_quantity"], r["quantity_difference"], r["source_sales"], r["database_sales"], r["sales_difference"]] for r in report["reconciliation"]],
        currency_columns=(8, 9, 10))
    add_export_sheet(workbook, "원본 미반영",
        ["플랫폼", "원본 행", "주문번호", "수량", "매출"],
        [[r["platform"], r["source_row"], r["order_number"], r["quantity"], r["sales_amount"]] for r in report["missing_source_rows"]], currency_columns=(5,))
    add_export_sheet(workbook, "필수값 누락",
        ["플랫폼", "원본 행", "주문번호", "누락 항목", "제품사이즈", "수량", "매출"],
        [[r["platform"], r["source_row"], r["order_number"], r["issue_detail"], dashboard_size_label(r["product_option"]), r["quantity"], r["sales_amount"]] for r in report["missing_fields"]], currency_columns=(7,))
    add_export_sheet(workbook, "0원 거래",
        ["플랫폼", "원본 행", "주문번호", "제품사이즈", "수량", "매출"],
        [[r["platform"], r["source_row"], r["order_number"], dashboard_size_label(r["product_option"]), r["quantity"], r["sales_amount"]] for r in report["zero_amount_rows"]], currency_columns=(6,))
    add_export_sheet(workbook, "확인 필요",
        ["플랫폼", "원본 행", "주문번호", "제품사이즈", "수량", "매출"],
        [[r["platform"], r["source_row"], r["order_number"], dashboard_size_label(r["product_option"]), r["quantity"], r["sales_amount"]] for r in report["review_issue_rows"]], currency_columns=(6,))
    add_export_sheet(workbook, "중복 의심",
        ["플랫폼", "주문번호", "구매자", "제품사이즈", "수량", "매출", "행 수", "원본 행"],
        [[r["platform"], r["order_number"], r["buyer_name"], dashboard_size_label(r["product_option"]), r["quantity"], r["sales_amount"], r["record_count"], r["source_rows"]] for r in report["duplicate_suspects"]], currency_columns=(6,))
    return workbook_response(workbook, dated_excel_filename("매출대사"))


@app.get("/sales")
def sales_list(
    request: Request,
    platform: str = Query(default=""),
    keyword: str = Query(default=""),
    sales_month: str = Query(default=""),
    item_type: str = Query(default=""),
    needs_review: bool = Query(default=False),
    sort_by: str = Query(default="tax_date"),
    sort_dir: str = Query(default="desc"),
    page: int = Query(default=1, ge=1),
):
    conditions = ["is_active = 1", "batch_is_active = 1"]
    parameters: list[object] = []
    selected_month = optional_number(sales_month, 1, 12)

    if platform:
        conditions.append("platform = ?")
        parameters.append(platform)

    if keyword:
        conditions.append(
            "(order_number LIKE ? OR buyer_name LIKE ? OR product_option LIKE ?)"
        )
        search_value = f"%{keyword}%"
        parameters.extend(
            [
                search_value,
                search_value,
                search_value,
            ]
        )

    if selected_month is not None:
        conditions.append("sales_month = ?")
        parameters.append(selected_month)

    if item_type:
        add_product_type_condition(conditions, parameters, item_type)

    if needs_review:
        conditions.append("review_status = 'issue'")

    where_clause = ""

    if conditions:
        where_clause = "WHERE " + " AND ".join(conditions)

    sortable_columns = {
        "platform": "platform",
        "order_number": "order_number",
        "tax_date": "tax_date",
        "product_option": "product_option",
        "quantity": "quantity",
        "sales_amount": (
            "credit_card_amount + cash_receipt_amount "
            "+ tax_invoice_amount + other_amount"
        ),
        "fee_amount": "fee_amount_1 + fee_amount_2",
        "settlement_amount": "settlement_amount",
        "buyer_name": "buyer_name",
    }
    selected_sort = sort_by if sort_by in sortable_columns else "tax_date"
    selected_direction = "asc" if sort_dir.lower() == "asc" else "desc"
    order_expression = sortable_columns[selected_sort]

    if selected_sort == "order_number":
        order_clause = f"""
            CASE
                WHEN TRIM(order_number) GLOB '[0-9]*' THEN 0
                ELSE 1
            END ASC,
            CAST(TRIM(order_number) AS INTEGER) {selected_direction},
            TRIM(order_number) {selected_direction},
            id DESC
        """
    else:
        order_clause = (
            f"{order_expression} {selected_direction}, id DESC"
        )

    with get_connection() as connection:
        total_count = connection.execute(
            f"SELECT COUNT(*) FROM sales_records {where_clause}",
            parameters,
        ).fetchone()[0]
        summary = connection.execute(
            f"""
            SELECT
                COALESCE(SUM({SALE_QUANTITY_EXPRESSION}), 0) AS quantity,
                COALESCE(SUM({SALES_EXPRESSION}), 0) AS sales_amount,
                COALESCE(SUM(credit_card_amount), 0) AS credit_card_amount,
                COALESCE(SUM(cash_receipt_amount), 0) AS cash_receipt_amount,
                COALESCE(SUM(tax_invoice_amount), 0) AS tax_invoice_amount,
                COALESCE(SUM(other_amount), 0) AS other_amount
            FROM sales_records
            {where_clause}
            """,
            parameters,
        ).fetchone()
        page_size = 50
        total_pages = max(1, math.ceil(total_count / page_size))
        current_page = min(page, total_pages)

        records = connection.execute(
            f"""
            SELECT
                *,
                (
                    credit_card_amount
                    + cash_receipt_amount
                    + tax_invoice_amount
                    + other_amount
                ) AS sales_amount
            FROM sales_records
            {where_clause}
            ORDER BY {order_clause}
            LIMIT ? OFFSET ?
            """,
            [
                *parameters,
                page_size,
                (current_page - 1) * page_size,
            ],
        ).fetchall()

        platforms = connection.execute(
            """
            SELECT DISTINCT platform
            FROM sales_records
            ORDER BY platform
            """
        ).fetchall()

    return templates.TemplateResponse(
        request=request,
        name="sales.html",
        context={
            "records": records,
            "platforms": platforms,
            "selected_platform": platform,
            "keyword": keyword,
            "selected_month": selected_month,
            "selected_item_type": item_type,
            "needs_review": needs_review,
            "selected_sort": selected_sort,
            "selected_direction": selected_direction,
            "current_page": current_page,
            "total_pages": total_pages,
            "total_count": total_count,
            "summary": summary,
        },
    )


@app.get("/sales/{record_id}")
def sales_detail(request: Request, record_id: int):
    with get_connection() as connection:
        record = connection.execute(
            "SELECT * FROM sales_records WHERE id = ?",
            (record_id,),
        ).fetchone()

    if record is None:
        return RedirectResponse(url="/sales", status_code=303)

    return templates.TemplateResponse(
        request=request,
        name="sales_detail.html",
        context={"record": record},
    )


@app.get("/sales/{record_id}/edit")
def sales_edit(request: Request, record_id: int):
    with get_connection() as connection:
        record = connection.execute(
            "SELECT * FROM sales_records WHERE id = ?",
            (record_id,),
        ).fetchone()

    if record is None:
        return RedirectResponse(url="/sales", status_code=303)

    return templates.TemplateResponse(
        request=request,
        name="sales_edit.html",
        context={"record": record},
    )


@app.post("/sales/{record_id}/edit")
def update_sales_record(
    record_id: int,
    item_type: str = Form(default="제품"),
    product_option: str = Form(default=""),
    quantity: float = Form(default=0),
    credit_card_amount: float = Form(default=0),
    cash_receipt_amount: float = Form(default=0),
    tax_invoice_amount: float = Form(default=0),
    other_amount: float = Form(default=0),
    fee_amount_1: float = Form(default=0),
    fee_amount_2: float = Form(default=0),
    settlement_amount: float = Form(default=0),
    buyer_name: str = Form(default=""),
    note: str = Form(default=""),
    review_status: str = Form(default="normal"),
    review_note: str = Form(default=""),
):
    if review_status not in ("normal", "reviewed", "issue"):
        review_status = "normal"
    item_type = normalized_product_type(item_type)
    if item_type not in PRODUCT_TYPE_EQUIVALENTS:
        item_type = "제품"

    with get_connection() as connection:
        connection.execute(
            """
            UPDATE sales_records
            SET
                item_type = ?, product_option = ?, quantity = ?,
                credit_card_amount = ?, cash_receipt_amount = ?,
                tax_invoice_amount = ?, other_amount = ?,
                fee_amount_1 = ?, fee_amount_2 = ?,
                settlement_amount = ?, buyer_name = ?, note = ?,
                review_status = ?, review_note = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (
                item_type,
                product_option.strip() or None,
                quantity,
                credit_card_amount,
                cash_receipt_amount,
                tax_invoice_amount,
                other_amount,
                fee_amount_1,
                fee_amount_2,
                settlement_amount,
                buyer_name.strip() or None,
                note.strip() or None,
                review_status,
                review_note.strip() or None,
                record_id,
            ),
        )
        connection.execute(
            """
            INSERT INTO audit_logs (
                action, target_type, target_id, description
            ) VALUES ('update', 'sales_record', ?, '판매내역 수정 및 검수')
            """,
            (record_id,),
        )

    return RedirectResponse(url=f"/sales/{record_id}", status_code=303)


@app.post("/sales/{record_id}/exclude")
def exclude_sales_record(record_id: int):
    with get_connection() as connection:
        connection.execute(
            """
            UPDATE sales_records
            SET is_active = 0, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (record_id,),
        )
        connection.execute(
            """
            INSERT INTO audit_logs (
                action, target_type, target_id, description
            ) VALUES ('exclude', 'sales_record', ?, '판매내역 집계 제외')
            """,
            (record_id,),
        )

    return RedirectResponse(url="/sales", status_code=303)


@app.get("/exports/sales.xlsx")
def export_sales(
    platform: str = Query(default=""),
    sales_month: str = Query(default=""),
    item_type: str = Query(default=""),
):
    conditions = ["is_active = 1", "batch_is_active = 1"]
    parameters: list[object] = []
    selected_month = optional_number(sales_month, 1, 12)

    for column, value in (
        ("platform", platform),
        ("sales_month", selected_month),
    ):
        if value not in (None, ""):
            conditions.append(f"{column} = ?")
            parameters.append(value)

    if item_type:
        add_product_type_condition(conditions, parameters, item_type)

    where_clause = ""

    if conditions:
        where_clause = "WHERE " + " AND ".join(conditions)

    with get_connection() as connection:
        rows = connection.execute(
            f"SELECT * FROM sales_records {where_clause} ORDER BY id",
            parameters,
        ).fetchall()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "통합 판매내역"
    headers = [
        "플랫폼", "주문번호", "매출월", "신고일", "제품유형",
        "제품코드", "수량", "신용카드", "현금영수증",
        "세금계산서", "기타", "수수료1", "수수료2",
        "정산금액", "구매자", "제품사이즈", "비고",
    ]
    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B5B")

    for row in rows:
        worksheet.append(
            [
                row["platform"], row["order_number"], row["sales_month"],
                excel_date_value(row["tax_date"]),
                product_type_label(row["item_type"]),
                product_code_for_size(row["product_option"], row["item_type"]),
                row["quantity"],
                row["credit_card_amount"],
                row["cash_receipt_amount"], row["tax_invoice_amount"],
                row["other_amount"], row["fee_amount_1"],
                row["fee_amount_2"], row["settlement_amount"],
                row["buyer_name"], dashboard_size_label(row["product_option"]), row["note"],
            ]
        )

        worksheet.cell(worksheet.max_row, 4).number_format = "yyyy-mm-dd"

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column in worksheet.columns:
        width = min(
            max(len(str(cell.value or "")) for cell in column) + 2,
            35,
        )
        worksheet.column_dimensions[column[0].column_letter].width = width

    return workbook_response(workbook, dated_excel_filename("통합판매내역"))


@app.get("/imports")
def import_history(request: Request):
    with get_connection() as connection:
        batches = connection.execute(
            """
            SELECT
                import_batches.*,
                COUNT(import_issues.id) AS issue_count,
                (SELECT COUNT(*) FROM sales_records WHERE import_batch_id = import_batches.id) AS affected_sales_count,
                COALESCE((SELECT SUM(credit_card_amount + cash_receipt_amount + tax_invoice_amount + other_amount) FROM sales_records WHERE import_batch_id = import_batches.id), 0) AS affected_sales_amount,
                (SELECT COUNT(*) FROM hometax_records WHERE import_batch_id = import_batches.id) AS affected_hometax_count
            FROM import_batches
            LEFT JOIN import_issues
                ON import_issues.import_batch_id = import_batches.id
            GROUP BY import_batches.id
            ORDER BY import_batches.id DESC
            """
        ).fetchall()

    return templates.TemplateResponse(
        request=request,
        name="imports.html",
        context={
            "batches": batches,
        },
    )


PRODUCT_SUPPLY_PRICE = 500_000


def round_won(value: object) -> int:
    return int(
        Decimal(str(value or 0)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    )


def build_accounting_summary(row) -> dict[str, object]:
    total_sales = round_won(row["total_sales"])
    ancillary_expense = -round_won(
        Decimal(str(row["ancillary_sales"] or 0)) / Decimal("1.1")
    )
    delivery_expense = -round_won(
        Decimal(str(row["delivery_sales"] or 0)) / Decimal("1.1")
    )
    confirmed_advance = round_won(
        Decimal(str(row["confirmed_quantity"] or 0)) * PRODUCT_SUPPLY_PRICE
    )
    unspecified_advance = -round_won(
        Decimal(str(row["unspecified_sales"] or 0)) / Decimal("1.1")
    )

    debit_rows = [
        {"source": "전체 매출", "account": "외상매출금", "amount": total_sales},
        {"source": "부가제품", "account": "오므론", "amount": ancillary_expense},
        {"source": "배송비", "account": "배송비", "amount": delivery_expense},
        {"source": "제품확정", "account": "선수금", "amount": confirmed_advance},
        {"source": "사이즈미정", "account": "선수금", "amount": unspecified_advance},
    ]

    credit_rows = [
        {
            "source": "제품",
            "account": "제품매출",
            "amount": round_won(Decimal(str(row["product_quantity"] or 0)) * PRODUCT_SUPPLY_PRICE),
        },
        {
            "source": "제품확정",
            "account": "제품매출",
            "amount": round_won(Decimal(str(row["confirmed_quantity"] or 0)) * PRODUCT_SUPPLY_PRICE),
        },
        {
            "source": "제품취소",
            "account": "제품매출(취소)",
            "amount": round_won(Decimal(str(row["cancel_quantity"] or 0)) * PRODUCT_SUPPLY_PRICE),
        },
    ]

    payment_rows = []
    for label, key in (
        ("신용카드", "credit_card_amount"),
        ("현금영수증", "cash_receipt_amount"),
        ("기타", "other_amount"),
        ("세금계산서", "tax_invoice_amount"),
    ):
        gross_amount = round_won(row[key])
        supply_amount = round_won(Decimal(str(gross_amount)) / Decimal("1.1"))
        vat_amount = round_won(Decimal(str(supply_amount)) * Decimal("0.1"))
        payment_rows.append(
            {
                "source": label,
                "account": "부가세예수금",
                "gross_amount": gross_amount,
                "supply_amount": supply_amount,
                "vat_amount": vat_amount,
            }
        )

    debit_total = sum(item["amount"] for item in debit_rows)
    credit_total = sum(item["amount"] for item in credit_rows) + sum(
        item["vat_amount"] for item in payment_rows
    )
    return {
        "debit_rows": debit_rows,
        "credit_rows": credit_rows,
        "payment_rows": payment_rows,
        "debit_total": debit_total,
        "credit_total": credit_total,
        "difference": debit_total - credit_total,
    }


def build_accounting_difference_rows(records, overall_difference: int) -> list[dict[str, object]]:
    differences = []
    row_difference_total = 0
    for record in records:
        item_type = normalized_product_type(record["item_type"])
        gross_amount = sum(
            record[key] or 0
            for key in (
                "credit_card_amount", "cash_receipt_amount",
                "tax_invoice_amount", "other_amount",
            )
        )
        row = {
            "total_sales": gross_amount,
            "ancillary_sales": gross_amount if item_type == "부가제품" else 0,
            "delivery_sales": gross_amount if item_type == "배송비" else 0,
            "unspecified_sales": gross_amount if item_type == "사이즈미정" else 0,
            "product_quantity": record["quantity"] if item_type == "제품" else 0,
            "confirmed_quantity": record["quantity"] if item_type == "제품확정" else 0,
            "cancel_quantity": -abs(record["quantity"] or 0) if item_type == "제품취소" else 0,
            "credit_card_amount": record["credit_card_amount"],
            "cash_receipt_amount": record["cash_receipt_amount"],
            "tax_invoice_amount": record["tax_invoice_amount"],
            "other_amount": record["other_amount"],
        }
        result = build_accounting_summary(row)
        difference = result["difference"]
        row_difference_total += difference
        if difference:
            differences.append({
                "record_id": record["id"],
                "order_number": record["order_number"],
                "item_type": item_type,
                "product_option": record["product_option"] or "-",
                "debit_amount": result["debit_total"],
                "credit_amount": result["credit_total"],
                "difference": difference,
                "note": record["note"] or "-",
            })

    rounding_difference = overall_difference - row_difference_total
    if rounding_difference:
        differences.append({
            "record_id": None,
            "order_number": "전체 반올림 조정",
            "item_type": "반올림",
            "product_option": "-",
            "debit_amount": 0,
            "credit_amount": 0,
            "difference": rounding_difference,
            "note": "결제수단별 합계 후 원 단위 반올림으로 발생한 차이",
        })
    return differences


def load_accounting_data(connection, sales_month: int | None, platform: str) -> dict[str, object]:
    month_rows = connection.execute(
        """SELECT DISTINCT sales_month FROM sales_records
        WHERE sales_month IS NOT NULL AND is_active = 1 AND batch_is_active = 1
        ORDER BY sales_month DESC"""
    ).fetchall()
    months = [item["sales_month"] for item in month_rows]
    selected_month = sales_month if sales_month in months else (months[0] if months else None)

    platform_rows = connection.execute(
        """SELECT DISTINCT platform FROM sales_records
        WHERE sales_month = ? AND is_active = 1 AND batch_is_active = 1
        ORDER BY platform""",
        (selected_month,),
    ).fetchall() if selected_month is not None else []
    platform_order = {name: index for index, name in enumerate(SUPPORTED_PLATFORMS)}
    platforms = sorted(
        (item["platform"] for item in platform_rows),
        key=lambda name: (platform_order.get(name, len(platform_order)), name),
    )
    selected_platform = platform if platform in platforms else (platforms[0] if platforms else "")

    summary_row = connection.execute(
        f"""SELECT
            COALESCE(SUM({SALES_EXPRESSION}), 0) AS total_sales,
            COALESCE(SUM(CASE WHEN item_type IN ('부가제품', '부가상품') THEN {SALES_EXPRESSION} ELSE 0 END), 0) AS ancillary_sales,
            COALESCE(SUM(CASE WHEN item_type = '배송비' THEN {SALES_EXPRESSION} ELSE 0 END), 0) AS delivery_sales,
            COALESCE(SUM(CASE WHEN item_type = '사이즈미정' THEN {SALES_EXPRESSION} ELSE 0 END), 0) AS unspecified_sales,
            COALESCE(SUM(CASE WHEN item_type IN ('제품', '상품') THEN quantity ELSE 0 END), 0) AS product_quantity,
            COALESCE(SUM(CASE WHEN item_type = '제품확정' THEN quantity ELSE 0 END), 0) AS confirmed_quantity,
            COALESCE(SUM(CASE WHEN item_type IN ('제품취소', '취소', '본품취소') THEN -ABS(quantity) ELSE 0 END), 0) AS cancel_quantity,
            COALESCE(SUM(credit_card_amount), 0) AS credit_card_amount,
            COALESCE(SUM(cash_receipt_amount), 0) AS cash_receipt_amount,
            COALESCE(SUM(tax_invoice_amount), 0) AS tax_invoice_amount,
            COALESCE(SUM(other_amount), 0) AS other_amount
        FROM sales_records
        WHERE sales_month = ? AND platform = ?
          AND is_active = 1 AND batch_is_active = 1""",
        (selected_month, selected_platform),
    ).fetchone()
    summary = build_accounting_summary(summary_row)
    records = connection.execute(
        """SELECT id, order_number, item_type, product_option, quantity,
            credit_card_amount, cash_receipt_amount, tax_invoice_amount,
            other_amount, note
        FROM sales_records
        WHERE sales_month = ? AND platform = ?
          AND is_active = 1 AND batch_is_active = 1
        ORDER BY id""",
        (selected_month, selected_platform),
    ).fetchall()
    return {
        "months": months,
        "platforms": platforms,
        "selected_month": selected_month,
        "selected_platform": selected_platform,
        "product_supply_price": PRODUCT_SUPPLY_PRICE,
        "raw_totals": {
            "ancillary_sales": round_won(summary_row["ancillary_sales"]),
            "delivery_sales": round_won(summary_row["delivery_sales"]),
            "unspecified_sales": round_won(summary_row["unspecified_sales"]),
        },
        "difference_rows": build_accounting_difference_rows(
            records, summary["difference"]
        ),
        **summary,
    }


@app.get("/accounting")
def accounting(request: Request, sales_month: int | None = None, platform: str = ""):
    with get_connection() as connection:
        data = load_accounting_data(connection, sales_month, platform)
    return templates.TemplateResponse(
        request=request,
        name="accounting.html",
        context=data,
    )


@app.get("/exports/accounting.xlsx")
def export_accounting(sales_month: int | None = None, platform: str = ""):
    with get_connection() as connection:
        initial_data = load_accounting_data(connection, sales_month, platform)
        selected_month = initial_data["selected_month"]
        months = initial_data["months"]
        month_data = {}
        all_platforms = set()
        for month in months:
            month_index = load_accounting_data(connection, month, "")
            all_platforms.update(month_index["platforms"])
            for item in month_index["platforms"]:
                month_data[(month, item)] = load_accounting_data(
                    connection, month, item
                )
        platform_order = {name: index for index, name in enumerate(SUPPORTED_PLATFORMS)}
        platforms = sorted(
            all_platforms,
            key=lambda name: (platform_order.get(name, len(platform_order)), name),
        )

    template_path = APP_DIRECTORY / "assets" / "accounting_export_template.xlsx"
    workbook = load_workbook(template_path, data_only=False, keep_links=False)
    worksheet = workbook["Sheet2"]
    worksheet.title = "회계처리"
    worksheet["C2"] = f"{selected_month}월" if selected_month else "-"
    source_sheet = workbook.create_sheet("계산데이터")
    source_sheet.append([
        "기준월", "플랫폼", "외상매출금", "오므론", "오므론 공급대가",
        "배송비", "배송비 공급대가", "제품확정 선수금", "사이즈미정 선수금",
        "사이즈미정 공급대가", "제품매출", "제품수량", "제품확정 매출",
        "제품확정 수량", "제품취소 매출", "제품취소 수량",
        "신용카드 부가세", "신용카드 공급가액", "신용카드 공급대가",
        "현금영수증 부가세", "현금영수증 공급가액", "현금영수증 공급대가",
        "기타 부가세", "기타 공급가액", "기타 공급대가",
        "세금계산서 부가세", "세금계산서 공급가액", "세금계산서 공급대가",
        "수수료 차액",
    ])
    for month in months:
        for item in platforms:
            data = month_data.get((month, item))
            if not data:
                continue
            debit, credit, payments = data["debit_rows"], data["credit_rows"], data["payment_rows"]
            raw = data["raw_totals"]
            source_sheet.append([
                f"{month}월", item, debit[0]["amount"], debit[1]["amount"], -raw["ancillary_sales"],
                debit[2]["amount"], -raw["delivery_sales"], debit[3]["amount"], debit[4]["amount"],
                -raw["unspecified_sales"], credit[0]["amount"], round_won(credit[0]["amount"] / PRODUCT_SUPPLY_PRICE),
                credit[1]["amount"], round_won(credit[1]["amount"] / PRODUCT_SUPPLY_PRICE),
                credit[2]["amount"], round_won(credit[2]["amount"] / PRODUCT_SUPPLY_PRICE),
                *[value for payment in payments for value in (
                    payment["vat_amount"], payment["supply_amount"], payment["gross_amount"]
                )],
                -data["difference"],
            ])

    month_list_column = 31
    for row_number, month in enumerate(months, start=2):
        source_sheet.cell(row_number, month_list_column, f"{month}월")
    if months:
        worksheet.data_validations.dataValidation = []
        validation = DataValidation(
            type="list",
            formula1='"' + ",".join(f"{month}월" for month in months) + '"',
        )
        worksheet.add_data_validation(validation)
        validation.add(worksheet["C2"])
    for cell_reference in ("B3", "B18", "B33"):
        worksheet[cell_reference] = PRODUCT_SUPPLY_PRICE
        worksheet[cell_reference].number_format = "#,##0_);[Red](#,##0)"

    block_positions = (
        (5, 4), (5, 11), (5, 18),
        (20, 4), (20, 11), (20, 18),
        (35, 4), (35, 11), (35, 18),
    )
    money_format = "#,##0_);[Red](#,##0)"

    for index, (start_row, start_column) in enumerate(block_positions):
        for row_number in range(start_row, start_row + 13):
            for column_number in range(start_column, start_column + 6):
                worksheet.cell(row_number, column_number).value = None
        if index >= len(platforms):
            continue
        platform_name = platforms[index]
        rows = [
            (platform_name, "차변", "외상매출금", 3, None, None),
            (None, "차변", "오므론", 4, None, 5),
            (None, "차변", "배송비", 6, None, 7),
            (None, "차변", "선수금", 8, None, None),
            (None, "차변", "선수금", 9, None, 10),
            (None, "대변", "제품매출", 11, None, 12),
            (None, "대변", "제품매출", 13, None, 14),
            (None, "대변", "제품매출(취소)", 15, None, 16),
            (None, "대변", "부가세예수금", 17, 18, 19),
            (None, "대변", "부가세예수금", 20, 21, 22),
            (None, "대변", "부가세예수금", 23, 24, 25),
            (None, "대변", "부가세예수금", 26, 27, 28),
            (None, "차변", "수수료(차액)", 29, None, None),
        ]
        for row_offset, values in enumerate(rows):
            for column_offset, value in enumerate(values):
                cell = worksheet.cell(start_row + row_offset, start_column + column_offset)
                if isinstance(value, int) and column_offset >= 3:
                    source_column = source_sheet.cell(1, value).column_letter
                    platform_cell = worksheet.cell(start_row, start_column).coordinate
                    cell.value = (
                        f'=IFERROR(SUMIFS(\'계산데이터\'!${source_column}$2:${source_column}${source_sheet.max_row},'
                        f'\'계산데이터\'!$A$2:$A${source_sheet.max_row},$C$2,'
                        f'\'계산데이터\'!$B$2:$B${source_sheet.max_row},{platform_cell}),0)'
                    )
                else:
                    cell.value = value
                if column_offset >= 3:
                    cell.number_format = money_format

    for row_number in range(6, 17):
        for column_number in range(25, 29):
            worksheet.cell(row_number, column_number).value = None
    for index, platform_name in enumerate(platforms[:11], start=6):
        worksheet.cell(index, 25, platform_name)
    for column_number in range(26, 29):
        worksheet.cell(17, column_number, f"=SUM({worksheet.cell(6, column_number).coordinate}:{worksheet.cell(16, column_number).coordinate})")
        worksheet.cell(17, column_number).number_format = money_format

    detail_rows = []
    for month in months:
        for platform_name in platforms:
            data = month_data.get((month, platform_name))
            if not data:
                continue
            detail_rows.extend([
            [
                f"{month}월", data["selected_platform"], item["order_number"], item["item_type"],
                dashboard_size_label(item["product_option"]), item["debit_amount"],
                item["credit_amount"], item["difference"], item["note"],
            ]
            for item in data["difference_rows"]
            ])
    add_export_sheet(
        workbook,
        "차이 상세내역",
        ["기준월", "플랫폼", "주문번호", "제품유형", "제품사이즈", "차변", "대변", "차이", "비고"],
        detail_rows,
        currency_columns=(6, 7, 8),
    )
    detail_sheet = workbook["차이 상세내역"]
    for column, width in {
        "A": 10, "B": 14, "C": 20, "D": 14, "E": 18,
        "F": 16, "G": 16, "H": 16, "I": 55,
    }.items():
        detail_sheet.column_dimensions[column].width = width
    for row_number in range(2, detail_sheet.max_row + 1):
        detail_sheet.cell(row_number, 9).alignment = Alignment(
            vertical="center", wrap_text=True
        )
        detail_sheet.row_dimensions[row_number].height = 30
    detail_sheet.auto_filter.ref = detail_sheet.dimensions
    source_sheet.sheet_state = "hidden"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    return workbook_response(workbook, dated_excel_filename("회계처리"))


@app.get("/reports")
def reports(request: Request):
    with get_connection() as connection:
        monthly = connection.execute(
            f"""
            SELECT
                SUBSTR(tax_date, 1, 7) AS period,
                SUM(quantity) AS quantity,
                SUM({SALES_EXPRESSION}) AS sales_amount
            FROM sales_records
            WHERE tax_date IS NOT NULL AND is_active = 1 AND batch_is_active = 1
            GROUP BY period
            ORDER BY period
            """
        ).fetchall()
        products = connection.execute(
            f"""
            SELECT
                item_type,
                COALESCE(product_option, '미지정') AS product_option,
                SUM(quantity) AS quantity,
                SUM({SALES_EXPRESSION}) AS sales_amount
            FROM sales_records
            WHERE is_active = 1 AND batch_is_active = 1
            GROUP BY item_type, product_option
            ORDER BY quantity DESC
            LIMIT 100
            """
        ).fetchall()
        taxes = connection.execute(
            """
            SELECT
                SUBSTR(tax_date, 1, 7) AS period,
                SUM(credit_card_amount) AS credit_card_amount,
                SUM(cash_receipt_amount) AS cash_receipt_amount,
                SUM(tax_invoice_amount) AS tax_invoice_amount,
                SUM(other_amount) AS other_amount,
                ROUND(SUM(
                    credit_card_amount + cash_receipt_amount
                    + tax_invoice_amount + other_amount
                ) / 11.0) AS vat_amount
            FROM sales_records
            WHERE tax_date IS NOT NULL AND is_active = 1 AND batch_is_active = 1
            GROUP BY period
            ORDER BY period
            """
        ).fetchall()
        accounting = connection.execute(
            f"""
            SELECT
                platform,
                SUM({SALES_EXPRESSION}) AS sales_amount,
                ROUND(SUM({SALES_EXPRESSION}) / 1.1) AS supply_amount,
                ROUND(SUM({SALES_EXPRESSION}) / 11.0) AS vat_amount,
                SUM(fee_amount_1 + fee_amount_2) AS fee_amount
            FROM sales_records
            WHERE is_active = 1 AND batch_is_active = 1
            GROUP BY platform
            ORDER BY sales_amount DESC
            """
        ).fetchall()
        hometax_reconciliation = connection.execute(
            f"""
            WITH sales AS (
                SELECT platform, SUBSTR(tax_date, 1, 7) AS period,
                    SUM({SALES_EXPRESSION}) AS sales_amount
                FROM sales_records
                WHERE tax_date IS NOT NULL AND is_active = 1 AND batch_is_active = 1
                GROUP BY platform, period
            ),
            hometax AS (
                SELECT platform, SUBSTR(sales_date, 1, 7) AS period,
                    SUM(total_amount) AS hometax_amount
                FROM hometax_records
                WHERE platform IS NOT NULL
                    AND batch_is_active = 1
                    AND platform != '취소'
                    AND sales_date IS NOT NULL
                GROUP BY platform, period
            ),
            keys AS (
                SELECT platform, period FROM sales
                UNION
                SELECT platform, period FROM hometax
            )
            SELECT keys.period, keys.platform,
                COALESCE(sales.sales_amount, 0) AS sales_amount,
                COALESCE(hometax.hometax_amount, 0) AS hometax_amount,
                COALESCE(sales.sales_amount, 0)
                    - COALESCE(hometax.hometax_amount, 0) AS difference_amount
            FROM keys
            LEFT JOIN sales
                ON sales.platform = keys.platform
                AND sales.period = keys.period
            LEFT JOIN hometax
                ON hometax.platform = keys.platform
                AND hometax.period = keys.period
            ORDER BY keys.period, keys.platform
            """
        ).fetchall()

    return templates.TemplateResponse(
        request=request,
        name="reports.html",
        context={
            "monthly": monthly,
            "products": products,
            "taxes": taxes,
            "accounting": accounting,
            "hometax_reconciliation": hometax_reconciliation,
        },
    )


@app.get("/exports/reports.xlsx")
def export_reports():
    with get_connection() as connection:
        monthly = connection.execute(f"SELECT SUBSTR(tax_date,1,7) period, SUM(quantity) quantity, SUM({SALES_EXPRESSION}) sales_amount FROM sales_records WHERE tax_date IS NOT NULL AND is_active=1 AND batch_is_active=1 GROUP BY period ORDER BY period").fetchall()
        products = connection.execute(f"SELECT item_type, COALESCE(product_option,'미지정') product_option, SUM(quantity) quantity, SUM({SALES_EXPRESSION}) sales_amount FROM sales_records WHERE is_active=1 AND batch_is_active=1 GROUP BY item_type,product_option ORDER BY quantity DESC LIMIT 100").fetchall()
        taxes = connection.execute("SELECT SUBSTR(tax_date,1,7) period, SUM(credit_card_amount) credit_card_amount, SUM(cash_receipt_amount) cash_receipt_amount, SUM(tax_invoice_amount) tax_invoice_amount, SUM(other_amount) other_amount, ROUND(SUM(credit_card_amount+cash_receipt_amount+tax_invoice_amount+other_amount)/11.0) vat_amount FROM sales_records WHERE tax_date IS NOT NULL AND is_active=1 AND batch_is_active=1 GROUP BY period ORDER BY period").fetchall()
        accounting = connection.execute(f"SELECT platform, SUM({SALES_EXPRESSION}) sales_amount, ROUND(SUM({SALES_EXPRESSION})/1.1) supply_amount, ROUND(SUM({SALES_EXPRESSION})/11.0) vat_amount, SUM(fee_amount_1+fee_amount_2) fee_amount FROM sales_records WHERE is_active=1 AND batch_is_active=1 GROUP BY platform ORDER BY sales_amount DESC").fetchall()
        reconciliation = connection.execute(f"""WITH sales AS (SELECT platform,SUBSTR(tax_date,1,7) period,SUM({SALES_EXPRESSION}) sales_amount FROM sales_records WHERE tax_date IS NOT NULL AND is_active=1 AND batch_is_active=1 GROUP BY platform,period), hometax AS (SELECT platform,SUBSTR(sales_date,1,7) period,SUM(total_amount) hometax_amount FROM hometax_records WHERE platform IS NOT NULL AND batch_is_active=1 AND platform!='취소' AND sales_date IS NOT NULL GROUP BY platform,period), keys AS (SELECT platform,period FROM sales UNION SELECT platform,period FROM hometax) SELECT keys.period,keys.platform,COALESCE(sales.sales_amount,0) sales_amount,COALESCE(hometax.hometax_amount,0) hometax_amount,COALESCE(sales.sales_amount,0)-COALESCE(hometax.hometax_amount,0) difference_amount FROM keys LEFT JOIN sales ON sales.platform=keys.platform AND sales.period=keys.period LEFT JOIN hometax ON hometax.platform=keys.platform AND hometax.period=keys.period ORDER BY keys.period,keys.platform""").fetchall()

    workbook = Workbook()
    add_export_sheet(workbook, "월별 매출", ["연월", "수량", "매출"], [[r["period"], r["quantity"], r["sales_amount"]] for r in monthly], currency_columns=(3,))
    add_export_sheet(workbook, "제품 사이즈", ["유형", "제품사이즈", "수량", "매출"], [[product_type_label(r["item_type"]), dashboard_size_label(r["product_option"]), r["quantity"], r["sales_amount"]] for r in products], currency_columns=(4,))
    add_export_sheet(workbook, "부가세 기초자료", ["연월", "신용카드", "현금영수증", "세금계산서", "기타", "예상 부가세"], [[r["period"], r["credit_card_amount"], r["cash_receipt_amount"], r["tax_invoice_amount"], r["other_amount"], r["vat_amount"]] for r in taxes], currency_columns=(2,3,4,5,6))
    add_export_sheet(workbook, "홈택스 대사", ["연월", "플랫폼", "웹 판매자료", "홈택스 자료", "차이"], [[r["period"], r["platform"], r["sales_amount"], r["hometax_amount"], r["difference_amount"]] for r in reconciliation], currency_columns=(3,4,5))
    add_export_sheet(workbook, "회계처리 집계", ["플랫폼", "매출", "공급가액", "부가세", "수수료"], [[r["platform"], r["sales_amount"], r["supply_amount"], r["vat_amount"], r["fee_amount"]] for r in accounting], currency_columns=(2,3,4,5))
    return workbook_response(workbook, dated_excel_filename("분석세무회계"))


@app.get("/hometax")
def hometax_management(
    request: Request,
    quarter: str = Query(default=""),
):
    conditions = ["batch_is_active = 1"]
    parameters: list[object] = []
    selected_quarter = optional_number(quarter, 1, 4)

    if selected_quarter is not None:
        start_month = (selected_quarter - 1) * 3 + 1
        end_month = start_month + 2
        conditions.append("sales_month BETWEEN ? AND ?")
        parameters.extend([start_month, end_month])

    where_clause = ""

    if conditions:
        where_clause = "WHERE " + " AND ".join(conditions)

    with get_connection() as connection:
        records = connection.execute(
            f"""
            SELECT *
            FROM hometax_records
            {where_clause}
            ORDER BY sales_date DESC, id DESC
            """,
            parameters,
        ).fetchall()

    return templates.TemplateResponse(
        request=request,
        name="hometax.html",
        context={
            "records": records,
            "platforms": SUPPORTED_PLATFORMS,
            "selected_quarter": selected_quarter,
        },
    )


@app.get("/exports/hometax.xlsx")
def export_hometax(quarter: str = Query(default="")):
    conditions = ["batch_is_active = 1"]
    parameters: list[object] = []
    selected_quarter = optional_number(quarter, 1, 4)

    if selected_quarter is not None:
        start_month = (selected_quarter - 1) * 3 + 1
        conditions.append("sales_month BETWEEN ? AND ?")
        parameters.extend([start_month, start_month + 2])

    with get_connection() as connection:
        records = connection.execute(
            f"SELECT * FROM hometax_records WHERE {' AND '.join(conditions)} ORDER BY sales_date DESC,id DESC",
            parameters,
        ).fetchall()

    workbook = Workbook()
    add_export_sheet(
        workbook, "홈택스 플랫폼 지정",
        ["구분", "매출월", "매출일", "거래건수", "공급가액", "부가세", "봉사료", "합계금액", "신용카드", "기타결제", "승인번호", "판매대행사", "원본 플랫폼", "지정 플랫폼", "지정 상태", "비고"],
        [["현금과세" if r["record_type"] == "cash_taxable" else "신용카드·건별", r["sales_month"], r["sales_date"], r["transaction_count"], r["supply_amount"], r["vat_amount"], r["service_amount"], r["total_amount"], r["credit_card_amount"], r["other_payment_amount"], r["approval_number"], r["merchant_name"], r["source_platform"], r["platform"], {"manual": "직접 지정", "unassigned": "미지정"}.get(r["platform_status"], r["platform_status"]), r["note"]] for r in records],
        currency_columns=(5,6,7,8,9,10), date_columns=(3,),
    )
    label = f"홈택스플랫폼지정_{selected_quarter}분기" if selected_quarter else "홈택스플랫폼지정_전체"
    return workbook_response(workbook, dated_excel_filename(label))


@app.post("/hometax/{record_id}/platform")
def update_hometax_platform(
    record_id: int,
    platform: str = Form(...),
    quarter: int | None = Form(default=None),
):
    normalized_platform = platform.strip()

    if normalized_platform not in SUPPORTED_PLATFORMS:
        return RedirectResponse(url="/hometax", status_code=303)

    with get_connection() as connection:
        connection.execute(
            """
            UPDATE hometax_records
            SET platform = ?, platform_status = 'manual'
            WHERE id = ? AND batch_is_active = 1
            """,
            (normalized_platform, record_id),
        )
        connection.execute(
            """
            INSERT INTO audit_logs (
                action, target_type, target_id, description
            ) VALUES ('update', 'hometax_record', ?, ?)
            """,
            (record_id, f"홈택스 플랫폼 지정: {normalized_platform}"),
        )

    redirect_url = "/hometax"

    if quarter is not None and 1 <= quarter <= 4:
        redirect_url = f"/hometax?quarter={quarter}"

    return RedirectResponse(url=redirect_url, status_code=303)


@app.post("/imports/{batch_id}/cancel")
def cancel_import(batch_id: int):
    with get_connection() as connection:
        batch = connection.execute(
            "SELECT * FROM import_batches WHERE id = ?",
            (batch_id,),
        ).fetchone()
        if batch is None or batch["status"] != "completed":
            return RedirectResponse(url="/imports", status_code=303)
        type_updates = connection.execute(
            """
            SELECT * FROM import_record_updates
            WHERE import_batch_id = ? AND is_active = 1
            """,
            (batch_id,),
        ).fetchall()
        for update in type_updates:
            connection.execute(
                """
                UPDATE sales_records SET item_type = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ? AND item_type = ?
                """,
                (
                    update["old_item_type"], update["sales_record_id"],
                    update["new_item_type"],
                ),
            )
        connection.execute(
            "UPDATE import_record_updates SET is_active = 0 WHERE import_batch_id = ?",
            (batch_id,),
        )
        connection.execute(
            "UPDATE sales_records SET batch_is_active = 0 WHERE import_batch_id = ?",
            (batch_id,),
        )
        connection.execute(
            "UPDATE hometax_records SET batch_is_active = 0 WHERE import_batch_id = ?",
            (batch_id,),
        )
        connection.execute(
            "UPDATE import_batches SET status = 'cancelled', cancelled_at = CURRENT_TIMESTAMP WHERE id = ?",
            (batch_id,),
        )
        connection.execute(
            """
            INSERT INTO audit_logs (
                action, target_type, target_id, description
            ) VALUES ('cancel', 'import_batch', ?, ?)
            """,
            (
                batch_id,
                f"업로드 반영 취소: {batch['original_filename']}",
            ),
        )

    return RedirectResponse(url="/imports", status_code=303)


@app.post("/imports/{batch_id}/restore")
def restore_import(batch_id: int):
    with get_connection() as connection:
        batch = connection.execute(
            "SELECT * FROM import_batches WHERE id = ?", (batch_id,)
        ).fetchone()
        if batch is None or batch["status"] != "cancelled":
            return RedirectResponse(url="/imports", status_code=303)
        type_updates = connection.execute(
            """
            SELECT * FROM import_record_updates
            WHERE import_batch_id = ? AND is_active = 0
            """,
            (batch_id,),
        ).fetchall()
        for update in type_updates:
            connection.execute(
                """
                UPDATE sales_records SET item_type = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ? AND item_type = ?
                """,
                (
                    update["new_item_type"], update["sales_record_id"],
                    update["old_item_type"],
                ),
            )
        connection.execute(
            "UPDATE import_record_updates SET is_active = 1 WHERE import_batch_id = ?",
            (batch_id,),
        )
        connection.execute(
            """
            UPDATE sales_records AS target
            SET batch_is_active = 1
            WHERE import_batch_id = ?
              AND NOT EXISTS (
                SELECT 1 FROM sales_records AS other
                WHERE other.id != target.id
                  AND other.batch_is_active = 1
                  AND SUBSTR(other.row_fingerprint, 1, 64)
                      = SUBSTR(target.row_fingerprint, 1, 64)
              )
            """,
            (batch_id,),
        )
        connection.execute(
            """
            UPDATE hometax_records AS target
            SET batch_is_active = 1
            WHERE import_batch_id = ?
              AND NOT EXISTS (
                SELECT 1 FROM hometax_records AS other
                WHERE other.id != target.id
                  AND other.batch_is_active = 1
                  AND SUBSTR(other.row_fingerprint, 1, 64)
                      = SUBSTR(target.row_fingerprint, 1, 64)
              )
            """,
            (batch_id,),
        )
        connection.execute(
            "UPDATE import_batches SET status = 'completed', restored_at = CURRENT_TIMESTAMP WHERE id = ?",
            (batch_id,),
        )
        connection.execute(
            "INSERT INTO audit_logs (action, target_type, target_id, description) VALUES ('restore', 'import_batch', ?, ?)",
            (batch_id, f"업로드 다시 복원: {batch['original_filename']}"),
        )
    return RedirectResponse(url="/imports", status_code=303)


@app.post("/imports/{batch_id}/delete")
def delete_import(batch_id: int):
    stored_filename = None
    with get_connection() as connection:
        batch = connection.execute(
            "SELECT * FROM import_batches WHERE id = ?", (batch_id,)
        ).fetchone()
        if batch is None or batch["status"] == "processing":
            return RedirectResponse(url="/imports", status_code=303)

        type_updates = connection.execute(
            """
            SELECT * FROM import_record_updates
            WHERE import_batch_id = ? AND is_active = 1
            """,
            (batch_id,),
        ).fetchall()
        for update in type_updates:
            connection.execute(
                """
                UPDATE sales_records SET item_type = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ? AND item_type = ?
                """,
                (
                    update["old_item_type"], update["sales_record_id"],
                    update["new_item_type"],
                ),
            )

        stored_filename = batch["stored_filename"]
        connection.execute("DELETE FROM import_batches WHERE id = ?", (batch_id,))

    if stored_filename:
        upload_root = UPLOAD_DIRECTORY.resolve()
        stored_path = (UPLOAD_DIRECTORY / stored_filename).resolve()
        if stored_path.parent == upload_root and stored_path.exists():
            try:
                stored_path.unlink()
            except OSError:
                pass

    return RedirectResponse(url="/imports", status_code=303)


@app.get("/imports/{batch_id}")
def import_detail(
    request: Request,
    batch_id: int,
):
    with get_connection() as connection:
        batch = connection.execute(
            """
            SELECT *
            FROM import_batches
            WHERE id = ?
            """,
            (batch_id,),
        ).fetchone()

        issues = connection.execute(
            """
            SELECT *
            FROM import_issues
            WHERE import_batch_id = ?
            ORDER BY platform, source_row, id
            """,
            (batch_id,),
        ).fetchall()
        blocking_issue_count = sum(
            1 for issue in issues if issue["severity"] == "error"
        )
        preview_rows = []

        if batch is not None and batch["status"] == "preview":
            stored_path = ensure_batch_source_file(connection, batch)
            connection.execute("SAVEPOINT detail_preview")
            try:
                preview_result = import_workbook(
                    connection,
                    stored_path,
                    batch_id,
                )
                preview_rows = preview_result.preview_rows
            finally:
                connection.execute("ROLLBACK TO detail_preview")
                connection.execute("RELEASE detail_preview")

    if batch is None:
        return RedirectResponse(
            url="/imports",
            status_code=303,
        )

    return templates.TemplateResponse(
        request=request,
        name="import_detail.html",
        context={
            "batch": batch,
            "issues": issues,
            "blocking_issue_count": blocking_issue_count,
            "preview_rows": preview_rows,
        },
    )


@app.post("/imports")
def upload_workbook(file: UploadFile = File(...)):
    original_filename = file.filename or "upload.xlsx"

    if Path(original_filename).suffix.lower() != ".xlsx":
        return RedirectResponse(
            url="/imports?error=xlsx",
            status_code=303,
        )

    stored_filename = f"{uuid4().hex}.xlsx"
    stored_path = UPLOAD_DIRECTORY / stored_filename

    source_content = file.file.read()
    stored_path.write_bytes(source_content)

    with get_connection() as connection:
        cursor = connection.execute(
            """
            INSERT INTO import_batches (
                original_filename,
                stored_filename,
                source_content,
                status
            ) VALUES (?, ?, ?, 'preview')
            RETURNING id
            """,
            (
                original_filename,
                stored_filename,
                source_content,
            ),
        )
        batch_id = cursor.fetchone()["id"]

        preview_active = False

        try:
            connection.execute("SAVEPOINT import_preview")
            preview_active = True
            result = import_workbook(
                connection=connection,
                file_path=stored_path,
                batch_id=batch_id,
            )
            preview_issues = [
                dict(row)
                for row in connection.execute(
                    "SELECT * FROM import_issues WHERE import_batch_id = ?",
                    (batch_id,),
                ).fetchall()
            ]
            connection.execute("ROLLBACK TO import_preview")
            connection.execute("RELEASE import_preview")
            preview_active = False
            connection.executemany(
                """
                INSERT INTO import_issues (
                    import_batch_id, platform, source_row, order_number,
                    field_name, issue_code, message, severity
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        issue["import_batch_id"], issue["platform"],
                        issue["source_row"], issue["order_number"],
                        issue["field_name"], issue["issue_code"],
                        issue["message"], issue["severity"],
                    )
                    for issue in preview_issues
                ],
            )
            connection.execute(
                """
                UPDATE import_batches
                SET
                    status = 'preview',
                    total_rows = ?,
                    imported_rows = ?,
                    duplicate_rows = ?,
                    issue_rows = ?
                WHERE id = ?
                """,
                (
                    result.total_rows,
                    result.imported_rows,
                    result.duplicate_rows,
                    result.issue_rows,
                    batch_id,
                ),
            )
        except Exception as error:
            if preview_active:
                connection.execute("ROLLBACK TO import_preview")
                connection.execute("RELEASE import_preview")

            connection.execute(
                """
                UPDATE import_batches
                SET
                    status = 'failed',
                    error_message = ?
                WHERE id = ?
                """,
                (
                    str(error),
                    batch_id,
                ),
            )

    return RedirectResponse(
        url=f"/imports/{batch_id}",
        status_code=303,
    )


@app.post("/imports/{batch_id}/confirm")
def confirm_import(batch_id: int):
    with get_connection() as connection:
        batch = connection.execute(
            "SELECT * FROM import_batches WHERE id = ? AND status = 'preview'",
            (batch_id,),
        ).fetchone()

        if batch is None:
            return RedirectResponse(url="/imports", status_code=303)

        blocking_issue = connection.execute(
            """
            SELECT 1 FROM import_issues
            WHERE import_batch_id = ? AND severity = 'error'
            LIMIT 1
            """,
            (batch_id,),
        ).fetchone()

        if blocking_issue is not None:
            return RedirectResponse(
                url=f"/imports/{batch_id}?error=blocking",
                status_code=303,
            )

        stored_path = ensure_batch_source_file(connection, batch)

        try:
            connection.execute(
                "DELETE FROM import_issues WHERE import_batch_id = ?",
                (batch_id,),
            )
            result = import_workbook(
                connection=connection,
                file_path=stored_path,
                batch_id=batch_id,
            )
            connection.execute(
                """
                UPDATE import_batches
                SET status = 'completed', total_rows = ?,
                    imported_rows = ?, duplicate_rows = ?, issue_rows = ?
                WHERE id = ?
                """,
                (
                    result.total_rows,
                    result.imported_rows,
                    result.duplicate_rows,
                    result.issue_rows,
                    batch_id,
                ),
            )
            connection.execute(
                """
                INSERT INTO audit_logs (
                    action, target_type, target_id, description
                ) VALUES ('import', 'import_batch', ?, ?)
                """,
                (
                    batch_id,
                    f"엑셀 최종 반영: {batch['original_filename']} ({result.imported_rows}행)",
                ),
            )
        except Exception as error:
            connection.execute(
                """
                UPDATE import_batches
                SET status = 'failed', error_message = ?
                WHERE id = ?
                """,
                (str(error), batch_id),
            )

    return RedirectResponse(url=f"/imports/{batch_id}", status_code=303)
