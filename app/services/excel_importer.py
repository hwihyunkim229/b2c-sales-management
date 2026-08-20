from __future__ import annotations
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from hashlib import sha256
from pathlib import Path
import re
import sqlite3
from openpyxl import load_workbook

SUPPORTED_PLATFORMS = (
    "네이버",
    "쿠팡",
    "카페24",
    "11번가",
    "SSG",
    "지마켓",
    "옥션",
    "베네피아",
    "현대홈쇼핑",
    "현대백화점",
)

@dataclass
class ImportResult:
    total_rows: int = 0
    imported_rows: int = 0
    duplicate_rows: int = 0
    issue_rows: int = 0
    preview_rows: list[dict[str, object]] = field(default_factory=list, repr=False)

def _active_unique_fingerprint(
    connection: sqlite3.Connection,
    table_name: str,
    base_fingerprint: str,
    batch_id: int,
) -> str | None:
    rows = connection.execute(
        f"""
        SELECT row_fingerprint, batch_is_active
        FROM {table_name}
        WHERE row_fingerprint = ? OR row_fingerprint LIKE ?
        """,
        (base_fingerprint, f"{base_fingerprint}:%"),
    ).fetchall()

    if any(row["batch_is_active"] == 1 for row in rows):
        return None

    if not rows:
        return base_fingerprint

    return f"{base_fingerprint}:{batch_id}"

def _text(value: object) -> str:
    if value is None:
        return ""

    return str(value).strip()

def _number(value: object) -> float:
    if value in (None, ""):
        return 0.0

    if isinstance(value, (int, float)):
        return float(value)

    normalized = str(value).replace(",", "").strip()

    try:
        return float(normalized)
    except ValueError:
        return 0.0

def _sales_month(value: object, tax_date: object) -> int | None:
    match = re.search(r"(\d{1,2})", _text(value))

    if match:
        month = int(match.group(1))

        if 1 <= month <= 12:
            return month

    parsed_date = _date_value(tax_date)

    if parsed_date:
        return int(parsed_date[5:7])

    return None

def _date_value(value: object) -> str | None:
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")

    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, (int, float)):
        excel_epoch = datetime(1899, 12, 30)
        converted = excel_epoch + timedelta(days=float(value))

        if converted.time() == datetime.min.time():
            return converted.date().isoformat()

        return converted.isoformat(timespec="seconds")

    raw_value = _text(value)

    for date_format in (
        "%Y-%m-%d",
        "%Y.%m.%d",
        "%Y/%m/%d",
    ):
        try:
            return datetime.strptime(raw_value, date_format).date().isoformat()
        except ValueError:
            continue

    return raw_value

def _item_type(product_option: str, quantity: float) -> str:
    normalized = product_option.replace(" ", "").lower()

    if "취소" in normalized or quantity < 0:
        return "취소"

    if "배송비" in normalized:
        return "배송비"

    if "수수료" in normalized:
        return "수수료"

    if "오므론" in normalized:
        return "부가상품"

    return "상품"

def _fingerprint(
    platform: str,
    source_row: int,
    row_values: tuple[object, ...],
) -> str:
    normalized_values = [platform, str(source_row)]
    normalized_values.extend(_text(value) for value in row_values[:16])
    payload = "|".join(normalized_values)

    return sha256(payload.encode("utf-8")).hexdigest()

def _record_issue(
    connection: sqlite3.Connection,
    batch_id: int,
    platform: str,
    source_row: int,
    order_number: str,
    field_name: str,
    issue_code: str,
    message: str,
    severity: str = "warning",
) -> None:
    connection.execute(
        """
        INSERT INTO import_issues (
            import_batch_id,
            platform,
            source_row,
            order_number,
            field_name,
            issue_code,
            message,
            severity
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            batch_id,
            platform,
            source_row,
            order_number,
            field_name,
            issue_code,
            message,
            severity,
        ),
    )

def import_workbook(
    connection: sqlite3.Connection,
    file_path: Path,
    batch_id: int,
) -> ImportResult:
    workbook = load_workbook(
        filename=file_path,
        data_only=True,
        read_only=True,
    )
    result = ImportResult()

    for platform in SUPPORTED_PLATFORMS:
        if platform not in workbook.sheetnames:
            continue

        worksheet = workbook[platform]

        for source_row, row in enumerate(
            worksheet.iter_rows(
                min_row=4,
                values_only=True,
            ),
            start=4,
        ):
            order_number = _text(row[1] if len(row) > 1 else None)
            product_option = _text(row[4] if len(row) > 4 else None)
            quantity = _number(row[5] if len(row) > 5 else None)
            normalized_product = product_option.replace(" ", "").lower()
            preserve_without_order = (
                "배송비" in normalized_product
                or "오므론" in normalized_product
            )

            if not order_number and not preserve_without_order:
                has_sales_values = any(
                    value not in (None, "", 0)
                    for value in row[2:16]
                )

                if has_sales_values:
                    result.total_rows += 1
                    result.issue_rows += 1
                    _record_issue(
                        connection=connection,
                        batch_id=batch_id,
                        platform=platform,
                        source_row=source_row,
                        order_number="",
                        field_name="order_number",
                        issue_code="missing_order_number",
                        message="주문번호가 없어 반영할 수 없습니다. 원본 엑셀을 보완한 뒤 다시 업로드하세요.",
                        severity="error",
                    )
                    result.preview_rows.append(
                        {
                            "status": "blocked",
                            "platform": platform,
                            "source_row": source_row,
                            "order_number": "",
                            "tax_date": _date_value(row[3] if len(row) > 3 else None),
                            "product_option": _text(row[4] if len(row) > 4 else None),
                            "quantity": _number(row[5] if len(row) > 5 else None),
                            "sales_amount": sum(_number(row[index] if len(row) > index else None) for index in range(6, 10)),
                            "buyer_name": _text(row[14] if len(row) > 14 else None),
                        }
                    )
                continue

            result.total_rows += 1
            row_fingerprint = _fingerprint(
                platform,
                source_row,
                row,
            )
            stored_fingerprint = _active_unique_fingerprint(
                connection,
                "sales_records",
                row_fingerprint,
                batch_id,
            )

            if stored_fingerprint is None:
                result.duplicate_rows += 1
                result.preview_rows.append(
                    {
                        "status": "duplicate",
                        "platform": platform,
                        "source_row": source_row,
                        "order_number": order_number,
                        "tax_date": _date_value(row[3] if len(row) > 3 else None),
                        "product_option": product_option,
                        "quantity": quantity,
                        "sales_amount": sum(_number(row[index] if len(row) > index else None) for index in range(6, 10)),
                        "buyer_name": _text(row[14] if len(row) > 14 else None),
                    }
                )
                continue
            parsed_date = _date_value(row[3] if len(row) > 3 else None)
            parsed_month = _sales_month(
                row[2] if len(row) > 2 else None,
                row[3] if len(row) > 3 else None,
            )

            issues = []

            if not order_number:
                issues.append(
                    (
                        "order_number",
                        "missing_order_number_preserved",
                        "오므론·배송비 행으로 원본 보존했으며 매출 집계에서는 제외됩니다.",
                    )
                )

            if parsed_month is None:
                issues.append(
                    (
                        "sales_month",
                        "missing_sales_month",
                        "매출월을 확인할 수 없습니다.",
                    )
                )

            if parsed_date is None:
                issues.append(
                    (
                        "tax_date",
                        "missing_tax_date",
                        "세금신고기준 날짜가 없습니다.",
                    )
                )

            if not product_option:
                issues.append(
                    (
                        "product_option",
                        "missing_product_option",
                        "제품 또는 사이즈 값이 없습니다.",
                    )
                )

            if row[5] in (None, ""):
                issues.append(
                    (
                        "quantity",
                        "missing_quantity",
                        "판매수량이 없습니다.",
                    )
                )

            for field_name, issue_code, message in issues:
                _record_issue(
                    connection=connection,
                    batch_id=batch_id,
                    platform=platform,
                    source_row=source_row,
                    order_number=order_number,
                    field_name=field_name,
                    issue_code=issue_code,
                    message=message,
                )
                result.issue_rows += 1

            result.preview_rows.append(
                {
                    "status": "issue" if issues else "new",
                    "platform": platform,
                    "source_row": source_row,
                    "order_number": order_number,
                    "tax_date": parsed_date,
                    "product_option": product_option,
                    "quantity": quantity,
                    "sales_amount": sum(_number(row[index] if len(row) > index else None) for index in range(6, 10)),
                    "buyer_name": _text(row[14] if len(row) > 14 else None),
                }
            )

            insert_result = connection.execute(
                    """
                    INSERT INTO sales_records (
                        import_batch_id,
                        platform,
                        source_row,
                        order_number,
                        sales_month,
                        tax_date,
                        item_type,
                        product_option,
                        quantity,
                        credit_card_amount,
                        cash_receipt_amount,
                        tax_invoice_amount,
                        other_amount,
                        fee_amount_1,
                        fee_amount_2,
                        settlement_amount,
                        advance_payment,
                        buyer_name,
                        note,
                        review_status,
                        row_fingerprint
                    ) VALUES (
                        ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                        ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
                    )
                    ON CONFLICT (row_fingerprint) DO NOTHING
                    RETURNING id
                    """,
                    (
                        batch_id,
                        platform,
                        source_row,
                        order_number,
                        parsed_month,
                        parsed_date,
                        _item_type(product_option, quantity),
                        product_option or None,
                        quantity,
                        _number(row[6] if len(row) > 6 else None),
                        _number(row[7] if len(row) > 7 else None),
                        _number(row[8] if len(row) > 8 else None),
                        _number(row[9] if len(row) > 9 else None),
                        _number(row[10] if len(row) > 10 else None),
                        _number(row[11] if len(row) > 11 else None),
                        _number(row[12] if len(row) > 12 else None),
                        _number(row[13] if len(row) > 13 else None),
                        _text(row[14] if len(row) > 14 else None) or None,
                        _text(row[15] if len(row) > 15 else None) or None,
                        "issue" if issues else "normal",
                        stored_fingerprint,
                    ),
                )
            if insert_result.fetchone() is not None:
                result.imported_rows += 1
            else:
                result.duplicate_rows += 1

    _import_hometax_sheet(
        connection=connection,
        workbook=workbook,
        batch_id=batch_id,
    )

    workbook.close()
    return result

def _import_hometax_sheet(
    connection: sqlite3.Connection,
    workbook,
    batch_id: int,
) -> None:
    if "홈택스" not in workbook.sheetnames:
        return

    worksheet = workbook["홈택스"]

    for source_row in range(5, worksheet.max_row + 1):
        cash_values = tuple(
            worksheet.cell(source_row, column).value
            for column in range(2, 14)
        )
        cash_total = _number(cash_values[6])

        if cash_total or _text(cash_values[0]):
            fingerprint = sha256(
                ("hometax_cash|" + "|".join(_text(value) for value in cash_values)).encode("utf-8")
            ).hexdigest()
            fingerprint = _active_unique_fingerprint(
                connection, "hometax_records", fingerprint, batch_id
            )
            if fingerprint is None:
                continue
            connection.execute(
                """
                INSERT INTO hometax_records (
                    import_batch_id, record_type, platform, source_platform,
                    platform_status, sales_month,
                    sales_date, supply_amount, vat_amount, service_amount,
                    total_amount, approval_number, transaction_type, note,
                    row_fingerprint
                ) VALUES (?, 'cash_taxable', ?, ?, 'manual', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT (row_fingerprint) DO NOTHING
                """,
                (
                    batch_id,
                    _text(cash_values[0]) or None,
                    _text(cash_values[0]) or None,
                    _sales_month(cash_values[1], cash_values[2]),
                    _date_value(cash_values[2]),
                    _number(cash_values[3]),
                    _number(cash_values[4]),
                    _number(cash_values[5]),
                    cash_total,
                    _text(cash_values[7]) or None,
                    _text(cash_values[9]) or None,
                    _text(cash_values[11]) or None,
                    fingerprint,
                ),
            )

        card_values = tuple(
            worksheet.cell(source_row, column).value
            for column in range(15, 23)
        )
        card_total = _number(card_values[4])

        if card_total or _text(card_values[0]):
            fingerprint = sha256(
                ("hometax_card|" + "|".join(_text(value) for value in card_values)).encode("utf-8")
            ).hexdigest()
            fingerprint = _active_unique_fingerprint(
                connection, "hometax_records", fingerprint, batch_id
            )
            if fingerprint is None:
                continue
            connection.execute(
                """
                INSERT INTO hometax_records (
                    import_batch_id, record_type, platform, source_platform,
                    platform_status, sales_month,
                    sales_date, transaction_count, total_amount,
                    credit_card_amount, other_payment_amount, merchant_name,
                    row_fingerprint
                ) VALUES (?, 'card_batch', ?, ?, 'manual', ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT (row_fingerprint) DO NOTHING
                """,
                (
                    batch_id,
                    _text(card_values[0]) or None,
                    _text(card_values[0]) or None,
                    _sales_month(card_values[1], card_values[2]),
                    _date_value(card_values[2]),
                    int(_number(card_values[3])),
                    card_total,
                    _number(card_values[5]),
                    _number(card_values[6]),
                    _text(card_values[7]) or None,
                    fingerprint,
                ),
            )
