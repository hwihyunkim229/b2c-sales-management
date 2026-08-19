from collections import defaultdict
from pathlib import Path
import sqlite3
from openpyxl import load_workbook
from app.services.excel_importer import (
    SUPPORTED_PLATFORMS,
    _fingerprint,
    _number,
    _text,
)

_source_cache: dict[str, object] = {
    "signature": None,
    "records": {},
    "file_count": 0,
    "errors": [],
}

def _load_source_records(
    batches: list[sqlite3.Row],
    upload_directory: Path,
) -> tuple[dict[str, dict[str, object]], int, list[str]]:
    signature_parts = []

    for batch in batches:
        file_path = upload_directory / batch["stored_filename"]
        stat = file_path.stat() if file_path.exists() else None
        signature_parts.append(
            (
                batch["id"],
                batch["stored_filename"],
                stat.st_size if stat else None,
                stat.st_mtime_ns if stat else None,
            )
        )

    signature = tuple(signature_parts)

    if _source_cache["signature"] == signature:
        return (
            _source_cache["records"],
            _source_cache["file_count"],
            _source_cache["errors"],
        )

    source_records: dict[str, dict[str, object]] = {}
    source_file_count = 0
    source_errors: list[str] = []

    for batch in batches:
        file_path = upload_directory / batch["stored_filename"]

        if not file_path.exists():
            source_errors.append(
                f"{batch['original_filename']}: 원본 파일을 찾을 수 없습니다."
            )
            continue

        try:
            workbook = load_workbook(
                file_path, data_only=True, read_only=True
            )
        except Exception as error:
            source_errors.append(f"{batch['original_filename']}: {error}")
            continue

        source_file_count += 1

        try:
            for platform in SUPPORTED_PLATFORMS:
                if platform not in workbook.sheetnames:
                    continue

                worksheet = workbook[platform]

                for source_row, row in enumerate(
                    worksheet.iter_rows(min_row=4, values_only=True), start=4
                ):
                    order_number = _text(row[1] if len(row) > 1 else None)

                    if not order_number:
                        continue

                    fingerprint = _fingerprint(platform, source_row, row)
                    source_records[fingerprint] = {
                        "platform": platform,
                        "source_row": source_row,
                        "order_number": order_number,
                        "quantity": _number(row[5] if len(row) > 5 else None),
                        "sales_amount": sum(
                            _number(row[index] if len(row) > index else None)
                            for index in range(6, 10)
                        ),
                    }
        finally:
            workbook.close()

    _source_cache.update(
        {
            "signature": signature,
            "records": source_records,
            "file_count": source_file_count,
            "errors": source_errors,
        }
    )
    return source_records, source_file_count, source_errors

def build_validation_report(
    connection: sqlite3.Connection,
    upload_directory: Path,
) -> dict[str, object]:
    batches = connection.execute(
        """
        SELECT id, original_filename, stored_filename
        FROM import_batches
        WHERE status = 'completed'
        ORDER BY id
        """
    ).fetchall()
    source_records, source_file_count, source_errors = _load_source_records(
        batches, upload_directory
    )

    database_rows = connection.execute(
        """
        SELECT
            id,
            platform,
            source_row,
            order_number,
            sales_month,
            tax_date,
            product_option,
            quantity,
            buyer_name,
            review_status,
            is_active,
            batch_is_active,
            row_fingerprint,
            credit_card_amount
                + cash_receipt_amount
                + tax_invoice_amount
                + other_amount AS sales_amount
        FROM sales_records
        ORDER BY id
        """
    ).fetchall()
    active_rows = [
        row for row in database_rows
        if row["is_active"] == 1 and row["batch_is_active"] == 1
    ]
    database_by_fingerprint = {
        row["row_fingerprint"][:64]: row
        for row in active_rows
    }
    platform_values: dict[str, dict[str, float]] = defaultdict(
        lambda: {
            "source_count": 0,
            "source_quantity": 0,
            "source_sales": 0,
            "database_count": 0,
            "database_quantity": 0,
            "database_sales": 0,
        }
    )

    for record in source_records.values():
        values = platform_values[str(record["platform"])]
        values["source_count"] += 1
        values["source_quantity"] += float(record["quantity"])
        values["source_sales"] += float(record["sales_amount"])

    for row in active_rows:
        values = platform_values[row["platform"]]
        values["database_count"] += 1
        values["database_quantity"] += row["quantity"] or 0
        values["database_sales"] += row["sales_amount"] or 0

    reconciliation = []

    for platform in sorted(platform_values):
        values = platform_values[platform]
        reconciliation.append(
            {
                "platform": platform,
                **values,
                "count_difference": (
                    values["database_count"] - values["source_count"]
                ),
                "quantity_difference": (
                    values["database_quantity"] - values["source_quantity"]
                ),
                "sales_difference": (
                    values["database_sales"] - values["source_sales"]
                ),
            }
        )

    missing_source_rows = [
        record
        for fingerprint, record in source_records.items()
        if fingerprint not in database_by_fingerprint
    ]
    missing_fields = []
    zero_amount_rows = []
    review_issue_rows = []

    for row in active_rows:
        fields = []

        if row["sales_month"] is None:
            fields.append("매출월")

        if not row["tax_date"]:
            fields.append("일자")

        if not row["product_option"]:
            fields.append("상품/사이즈")

        if row["quantity"] is None:
            fields.append("수량")

        if fields:
            missing_fields.append(
                {
                    **dict(row),
                    "issue_detail": ", ".join(fields),
                }
            )

        if abs(row["sales_amount"] or 0) < 0.001:
            zero_amount_rows.append(dict(row))

        if row["review_status"] == "issue":
            review_issue_rows.append(dict(row))

    duplicate_groups: dict[tuple[object, ...], list[sqlite3.Row]] = defaultdict(list)

    for row in active_rows:
        duplicate_groups[
            (
                row["platform"],
                row["order_number"],
                row["product_option"],
                row["buyer_name"],
                row["quantity"],
                row["sales_amount"],
            )
        ].append(row)

    duplicate_suspects = []

    for rows in duplicate_groups.values():
        if len(rows) < 2:
            continue

        duplicate_suspects.append(
            {
                "platform": rows[0]["platform"],
                "order_number": rows[0]["order_number"],
                "product_option": rows[0]["product_option"],
                "buyer_name": rows[0]["buyer_name"],
                "quantity": rows[0]["quantity"],
                "sales_amount": rows[0]["sales_amount"],
                "record_count": len(rows),
                "source_rows": ", ".join(
                    str(row["source_row"])
                    for row in rows
                ),
                "record_id": rows[0]["id"],
            }
        )

    mismatch_count = sum(
        1
        for row in reconciliation
        if (
            row["count_difference"] != 0
            or abs(row["quantity_difference"]) > 0.001
            or abs(row["sales_difference"]) > 0.001
        )
    )
    issue_count = (
        len(missing_source_rows)
        + len(missing_fields)
        + len(zero_amount_rows)
        + len(review_issue_rows)
    )

    return {
        "source_file_count": source_file_count,
        "source_errors": source_errors,
        "active_count": len(active_rows),
        "excluded_count": sum(
            1 for row in database_rows
            if row["is_active"] == 0 and row["batch_is_active"] == 1
        ),
        "mismatch_count": mismatch_count,
        "issue_count": issue_count,
        "reconciliation": reconciliation,
        "missing_source_rows": missing_source_rows,
        "missing_fields": missing_fields,
        "zero_amount_rows": zero_amount_rows,
        "review_issue_rows": review_issue_rows,
        "duplicate_suspects": duplicate_suspects,
    }