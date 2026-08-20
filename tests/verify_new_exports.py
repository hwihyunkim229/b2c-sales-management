from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app


def check(path: str, expected_sheets: list[str]) -> None:
    response = TestClient(app).get(path)
    response.raise_for_status()
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == expected_sheets

    for worksheet in workbook.worksheets:
        assert worksheet.freeze_panes == "A2"
        assert worksheet.auto_filter.ref is not None
        assert all(cell.value for cell in worksheet[1])

    print(path, len(response.content), workbook.sheetnames)


def main() -> None:
    check("/exports/validation.xlsx", [
        "플랫폼 대사", "원본 미반영", "필수값 누락",
        "0원 거래", "확인 필요", "중복 의심",
    ])
    check("/exports/reports.xlsx", [
        "월별 매출", "제품 사이즈", "부가세 기초자료",
        "홈택스 대사", "회계처리 집계",
    ])
    check("/exports/hometax.xlsx?quarter=3", ["홈택스 플랫폼 지정"])


if __name__ == "__main__":
    main()
