from pathlib import Path
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import get_connection, initialize_database
from app.main import app
from app.services.excel_importer import import_workbook


def verify_import() -> None:
    initialize_database()

    connection = get_connection()
    connection.execute("BEGIN")

    try:
        cursor = connection.execute(
            """
            INSERT INTO import_batches (
                original_filename,
                stored_filename,
                status
            ) VALUES (?, ?, 'processing')
            """,
            (
                "validation.xlsx",
                "validation.xlsx",
            ),
        )
        result = import_workbook(
            connection=connection,
            file_path=Path("uploads/initial-source.xlsx"),
            batch_id=cursor.lastrowid,
        )

        assert result.total_rows == 671
        assert result.issue_rows >= 0

        print(result)
    finally:
        connection.rollback()
        connection.close()


def verify_pages() -> None:
    with TestClient(app) as client:
        for path in (
            "/",
            "/sales",
            "/validation",
            "/imports",
            "/imports/1",
        ):
            response = client.get(path)
            response.raise_for_status()

            print(path, response.status_code, len(response.content))

        export_response = client.get("/exports/sales.xlsx")
        export_response.raise_for_status()
        workbook = load_workbook(BytesIO(export_response.content))
        worksheet = workbook["통합 판매내역"]
        if worksheet.max_row > 1:
            assert worksheet["D2"].number_format == "yyyy-mm-dd"
            assert not isinstance(worksheet["D2"].value, str)


if __name__ == "__main__":
    verify_import()
    verify_pages()
