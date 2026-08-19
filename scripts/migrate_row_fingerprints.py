from pathlib import Path

from openpyxl import load_workbook

from app.database import get_connection
from app.services.excel_importer import (
    SUPPORTED_PLATFORMS,
    _fingerprint,
    _text,
    import_workbook,
)


SOURCE_PATH = Path("uploads/initial-source.xlsx")


def main() -> None:
    workbook = load_workbook(
        SOURCE_PATH,
        data_only=True,
        read_only=True,
    )
    fingerprints: dict[tuple[str, int], str] = {}

    for platform in SUPPORTED_PLATFORMS:
        if platform not in workbook.sheetnames:
            continue

        worksheet = workbook[platform]

        for source_row, row in enumerate(
            worksheet.iter_rows(min_row=4, values_only=True),
            start=4,
        ):
            order_number = _text(row[1] if len(row) > 1 else None)

            if not order_number:
                continue

            fingerprints[(platform, source_row)] = _fingerprint(
                platform,
                source_row,
                row,
            )

    with get_connection() as connection:
        batch = connection.execute(
            "SELECT id FROM import_batches ORDER BY id LIMIT 1"
        ).fetchone()

        if batch is None:
            raise RuntimeError("마이그레이션할 최초 업로드 배치가 없습니다.")

        batch_id = int(batch["id"])
        records = connection.execute(
            "SELECT id, platform, source_row FROM sales_records"
        ).fetchall()

        for record in records:
            key = (record["platform"], int(record["source_row"]))

            if key in fingerprints:
                connection.execute(
                    "UPDATE sales_records SET row_fingerprint = ? WHERE id = ?",
                    (f"migration-{record['id']}", record["id"]),
                )

        for record in records:
            key = (record["platform"], int(record["source_row"]))

            if key in fingerprints:
                connection.execute(
                    "UPDATE sales_records SET row_fingerprint = ? WHERE id = ?",
                    (fingerprints[key], record["id"]),
                )

        result = import_workbook(
            connection=connection,
            file_path=SOURCE_PATH,
            batch_id=batch_id,
        )
        record_count = connection.execute(
            "SELECT COUNT(*) AS count FROM sales_records WHERE import_batch_id = ?",
            (batch_id,),
        ).fetchone()["count"]
        connection.execute(
            """
            UPDATE import_batches
            SET
                total_rows = ?,
                imported_rows = ?,
                duplicate_rows = 0
            WHERE id = ?
            """,
            (result.total_rows, record_count, batch_id),
        )
        connection.execute(
            """
            INSERT INTO audit_logs (
                action,
                target_type,
                target_id,
                description
            ) VALUES (?, ?, ?, ?)
            """,
            (
                "dedupe_migration",
                "import_batch",
                batch_id,
                "원본 행 기준으로 중복 판정키를 변경하고 누락 행을 복원했습니다.",
            ),
        )

        total = connection.execute(
            """
            SELECT
                COUNT(*) AS count,
                SUM(
                    credit_card_amount
                    + cash_receipt_amount
                    + tax_invoice_amount
                    + other_amount
                ) AS sales_amount
            FROM sales_records
            WHERE is_active = 1
            """
        ).fetchone()

    print(
        {
            "source_rows": result.total_rows,
            "restored_rows": result.imported_rows,
            "existing_rows": result.duplicate_rows,
            "active_rows": total["count"],
            "sales_amount": total["sales_amount"],
        }
    )


if __name__ == "__main__":
    main()
