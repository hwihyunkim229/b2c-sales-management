from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from app.database import get_connection
from app.services.excel_importer import (
    SUPPORTED_PLATFORMS,
    _fingerprint,
    _number,
    _text,
)


SOURCE_PATH = Path("uploads/initial-source.xlsx")


def main() -> None:
    workbook = load_workbook(
        SOURCE_PATH,
        data_only=True,
        read_only=True,
    )
    print(
        "WORKBOOK CACHED TOTALS",
        {
            "dashboard_total_sales": workbook["Dashboard"]["A4"].value,
            "raw_sales_total": workbook["Raw_통합"]["L1"].value,
            "raw_supply_total": workbook["Raw_통합"]["M1"].value,
            "raw_vat_total": workbook["Raw_통합"]["N1"].value,
            "monthly_total": workbook["월별_매출내역"]["N31"].value,
        },
    )
    excel_totals = {}
    raw_totals = defaultdict(float)
    fingerprints = defaultdict(list)

    raw_sheet = workbook["Raw_통합"]

    for row in raw_sheet.iter_rows(min_row=5, values_only=True):
        platform = _text(row[0] if len(row) > 0 else None)

        if platform:
            raw_totals[platform] += _number(row[11] if len(row) > 11 else None)

    for platform in SUPPORTED_PLATFORMS:
        if platform not in workbook.sheetnames:
            continue

        worksheet = workbook[platform]
        total_amount = 0.0
        total_quantity = 0.0
        record_count = 0

        for source_row, row in enumerate(
            worksheet.iter_rows(min_row=4, values_only=True),
            start=4,
        ):
            order_number = _text(row[1] if len(row) > 1 else None)

            if not order_number:
                continue

            record_count += 1
            total_quantity += _number(row[5] if len(row) > 5 else None)
            total_amount += sum(
                _number(row[index] if len(row) > index else None)
                for index in range(6, 10)
            )
            fingerprint = _fingerprint(
                platform,
                source_row,
                row,
            )
            fingerprints[fingerprint].append(
                {
                    "platform": platform,
                    "source_row": source_row,
                    "order_number": order_number,
                    "product_option": _text(row[4] if len(row) > 4 else None),
                    "quantity": _number(row[5] if len(row) > 5 else None),
                    "sales_amount": sum(
                        _number(row[index] if len(row) > index else None)
                        for index in range(6, 10)
                    ),
                    "buyer": _text(row[14] if len(row) > 14 else None),
                }
            )

        excel_totals[platform] = {
            "records": record_count,
            "quantity": total_quantity,
            "sales_amount": total_amount,
        }

    with get_connection() as connection:
        database_rows = connection.execute(
            """
            SELECT
                platform,
                COUNT(*) AS records,
                SUM(quantity) AS quantity,
                SUM(
                    credit_card_amount
                    + cash_receipt_amount
                    + tax_invoice_amount
                    + other_amount
                ) AS sales_amount
            FROM sales_records
            WHERE is_active = 1
            GROUP BY platform
            ORDER BY platform
            """
        ).fetchall()

    database_totals = {
        row["platform"]: dict(row)
        for row in database_rows
    }

    print("PLATFORM RECONCILIATION")
    grand_excel = 0.0
    grand_database = 0.0

    for platform in SUPPORTED_PLATFORMS:
        excel = excel_totals.get(
            platform,
            {"records": 0, "quantity": 0, "sales_amount": 0},
        )
        database = database_totals.get(
            platform,
            {"records": 0, "quantity": 0, "sales_amount": 0},
        )
        difference = excel["sales_amount"] - database["sales_amount"]
        grand_excel += excel["sales_amount"]
        grand_database += database["sales_amount"]

        print(
            platform,
            {
                "excel_records": excel["records"],
                "db_records": database["records"],
                "excel_quantity": excel["quantity"],
                "db_quantity": database["quantity"],
                "excel_sales": excel["sales_amount"],
                "raw_sales": raw_totals.get(platform, 0),
                "db_sales": database["sales_amount"],
                "difference": difference,
                "raw_vs_db": raw_totals.get(platform, 0) - database["sales_amount"],
            },
        )

    print(
        "GRAND TOTAL",
        {
            "excel_sales": grand_excel,
            "db_sales": grand_database,
            "difference": grand_excel - grand_database,
        },
    )

    print("EXACT DUPLICATE SOURCE ROWS")

    for rows in fingerprints.values():
        if len(rows) > 1:
            print(rows)

    print("RAW ROW DIFFERENCES")

    for platform in SUPPORTED_PLATFORMS:
        if platform not in workbook.sheetnames:
            continue

        source_sheet = workbook[platform]
        source_by_row = {}

        for source_row, row in enumerate(
            source_sheet.iter_rows(min_row=4, values_only=True),
            start=4,
        ):
            if not _text(row[1] if len(row) > 1 else None):
                continue

            source_by_row[source_row] = sum(
                _number(row[index] if len(row) > index else None)
                for index in range(6, 10)
            )

        raw_by_row = defaultdict(float)

        for row in raw_sheet.iter_rows(min_row=5, values_only=True):
            if _text(row[0]) != platform:
                continue

            source_row = int(_number(row[1]))
            raw_by_row[source_row] += _number(row[11])

        for source_row in sorted(set(source_by_row) | set(raw_by_row)):
            difference = source_by_row.get(source_row, 0) - raw_by_row.get(source_row, 0)

            if abs(difference) > 0.001:
                print(
                    platform,
                    {
                        "source_row": source_row,
                        "source_sales": source_by_row.get(source_row, 0),
                        "raw_sales": raw_by_row.get(source_row, 0),
                        "difference": difference,
                    },
                )

    print("COUPANG SOURCE ROW 5 TRACE")
    coupang_row_5 = list(workbook["쿠팡"].iter_rows(min_row=5, max_row=5, values_only=True))[0]
    print(
        {
            "source_values": coupang_row_5[:16],
            "source_sales": sum(_number(coupang_row_5[index]) for index in range(6, 10)),
        }
    )

    for raw_row_number, row in enumerate(
        raw_sheet.iter_rows(min_row=5, values_only=True),
        start=5,
    ):
        if _text(row[0]) == "쿠팡" and int(_number(row[1])) == 5:
            print(
                {
                    "raw_row": raw_row_number,
                    "raw_values": row[:20],
                }
            )


if __name__ == "__main__":
    main()
